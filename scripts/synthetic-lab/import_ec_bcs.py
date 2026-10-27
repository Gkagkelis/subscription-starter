"""
Έρευνα Καταναλωτών ΕΕ (DG ECFIN Business & Consumer Surveys· στην Ελλάδα
τη διεξάγει το ΙΟΒΕ) -> οικονομικό κλίμα ανά ομάδα για το Synthetic Research Lab.

Πηγή: https://economy-finance.ec.europa.eu/economic-forecast-and-surveys/business-and-consumer-surveys/download-business-and-consumer-survey-data/time-series_en
Άδεια: επαναχρησιμοποίηση περιεχομένου της Ευρωπαϊκής Επιτροπής με αναφορά
πηγής (Απόφαση 2011/833/ΕΕ, CC BY 4.0) — επιτρέπεται και η αναδιανομή.

Οι τιμές είναι ΙΣΟΖΥΓΙΑ (balances, -100..+100: θετικές μείον αρνητικές
απαντήσεις), εποχικά διορθωμένα — ΟΧΙ ποσοστά απαντήσεων. Το Lab τα
χρησιμοποιεί ως πλαίσιο (οικονομικό κλίμα, διαφορές ομάδων) και ως τεκμήριο
στις παραδοχές· δεν τα μετατρέπει σε πιθανότητες απάντησης.

Χρήση (pyreadstat δεν χρειάζεται· απαιτεί openpyxl):
  # consumer_total_sa_nace2.zip και consumer_subsectors_nace2.zip αποσυμπιεσμένα στο --raw
  python scripts/synthetic-lab/import_ec_bcs.py --raw .synthetic-lab-data/raw/ec-bcs/unz \\
      --out lib/synthetic-lab/data/official/ec-bcs-consumer-greece.json --version 2609
"""

import argparse
import json
import os
from datetime import datetime, timezone

import openpyxl

QUESTIONS = {
    "COF": "Δείκτης καταναλωτικής εμπιστοσύνης",
    "1": "Οικονομική κατάσταση νοικοκυριού — τελευταίοι 12 μήνες",
    "2": "Οικονομική κατάσταση νοικοκυριού — επόμενοι 12 μήνες",
    "3": "Γενική οικονομική κατάσταση της χώρας — τελευταίοι 12 μήνες",
    "4": "Γενική οικονομική κατάσταση της χώρας — επόμενοι 12 μήνες",
    "5": "Τιμές — τελευταίοι 12 μήνες",
    "6": "Τιμές — επόμενοι 12 μήνες",
    "7": "Ανεργία — επόμενοι 12 μήνες",
    "8": "Είναι τώρα κατάλληλη στιγμή για μεγάλες αγορές;",
    "9": "Δαπάνες για μεγάλες αγορές — επόμενοι 12 μήνες",
    "11": "Αποταμίευση — επόμενοι 12 μήνες",
    "12": "Δήλωση για την οικονομική κατάσταση του νοικοκυριού (αποταμιεύει / χρεώνεται)",
}
GROUPS = {
    "TOT": ("all", "all"),
    "RE1": ("income_quartile", "1"), "RE2": ("income_quartile", "2"), "RE3": ("income_quartile", "3"), "RE4": ("income_quartile", "4"),
    "ED1": ("education_group", "low_education"), "ED2": ("education_group", "medium_education"), "ED3": ("education_group", "high_education"),
    "AG1": ("age_group", "16-29"), "AG2": ("age_group", "30-49"), "AG3": ("age_group", "50-64"), "AG4": ("age_group", "65+"),
    "MAL": ("gender", "male"), "FEM": ("gender", "female"),
}
KEEP_MONTHS = 36


def read_sheet(ws, prefix):
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    cols = {i: h for i, h in enumerate(header) if isinstance(h, str) and h.startswith(prefix)}
    out = {}
    for r in rows[1:]:
        d = r[0]
        if not isinstance(d, datetime):
            continue
        key = d.strftime("%Y-%m")
        for i, h in cols.items():
            v = r[i]
            if isinstance(v, (int, float)):
                out.setdefault(h, {})[key] = round(float(v), 1)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--version", default="")
    args = ap.parse_args()

    series = {}
    wb = openpyxl.load_workbook(os.path.join(args.raw, "consumer_total_sa_nace2.xlsx"), read_only=True)
    series.update(read_sheet(wb["CONSUMER MONTHLY"], "CONS.EL.TOT."))
    wb = openpyxl.load_workbook(os.path.join(args.raw, "consumer_subsectors_sa_m_nace2.xlsx"), read_only=True)
    for sheet in GROUPS:
        if sheet in wb.sheetnames:
            series.update(read_sheet(wb[sheet], f"CONS.EL.{sheet}."))

    months = sorted({m for s in series.values() for m in s})[-KEEP_MONTHS:]
    latest = months[-1]
    last12 = months[-12:]
    indicators = []
    for q, label in QUESTIONS.items():
        groups = {}
        for g, (dim, grp) in GROUPS.items():
            s = series.get(f"CONS.EL.{g}.{q}.BS.M")
            if not s:
                continue
            vals = [s[m] for m in last12 if m in s]
            groups.setdefault(dim, {})[grp] = {
                "latest": s.get(latest),
                "mean12m": round(sum(vals) / len(vals), 1) if vals else None,
                "series": {m: s[m] for m in months if m in s},
            }
        if groups:
            indicators.append({"id": f"bcs_{q.lower()}", "question": q, "label": label, "groups": groups})

    out = {
        "datasetId": f"ec-bcs-consumer-el-{latest.replace('-', '')}",
        "source": "Ευρωπαϊκή Επιτροπή, DG ECFIN — Έρευνα Καταναλωτών (Business and Consumer Surveys), Ελλάδα (ΙΟΒΕ)· εποχικά διορθωμένα ισοζύγια",
        "url": "https://economy-finance.ec.europa.eu/economic-forecast-and-surveys/business-and-consumer-surveys/download-business-and-consumer-survey-data/time-series_en",
        "license": "Ευρωπαϊκή Επιτροπή — επαναχρησιμοποίηση με αναφορά πηγής (Απόφαση 2011/833/ΕΕ, CC BY 4.0)",
        "fileVersion": args.version,
        "latestMonth": latest,
        "fetchedAt": datetime.now(timezone.utc).isoformat(),
        "unit": "ισοζύγιο (balance), -100..+100",
        "note": "Ισοζύγια = σταθμισμένες θετικές μείον αρνητικές απαντήσεις· ΔΕΝ είναι ποσοστά. Ομάδες: εισοδηματικά τεταρτημόρια νοικοκυριού, μόρφωση (πρωτοβάθμια/δευτεροβάθμια/ανώτερη), ηλικία 16-29/30-49/50-64/65+, φύλο.",
        "indicators": indicators,
    }
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    print(f"✓ {len(indicators)} δείκτες, τελευταίος μήνας {latest} -> {args.out}")
    for ind in indicators:
        if ind["question"] in ("COF", "8", "9", "11"):
            g = ind["groups"]
            print(" ", ind["label"], "| σύνολο", g.get("all", {}).get("all", {}).get("latest"),
                  "| εισόδημα", [g.get("income_quartile", {}).get(k, {}).get("latest") for k in "1234"])


if __name__ == "__main__":
    main()
