"""
PHASE6 — προβλέψεις για το τυφλό holdout H6 (προκαταχώριση: docs/synthetic-lab/PHASE6_PREREGISTRATION.md). ΜΙΑ εκτέλεση.
Διαβάζει ΜΟΝΟ visible/ (άλλες χώρες) και τα αρχεία Volume C των άλλων χωρών. ΔΕΝ διαβάζει locked/.
Γράφει predictions.json με sha256 (κλείδωμα) ΠΡΙΝ από τη βαθμολόγηση (gpm_h6_score.py).

  python3 -I scripts/synthetic-lab/persona/gpm_h6_predict.py --root .
"""
import argparse
import glob
import hashlib
import io
import json
import os
import re
import sys
import time
import zipfile

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

AGG = {"EU27", "EA", "EA20", "DEW", "DEE", "EL"}
DK_RE = re.compile(r"(don.?t know|dk\b|refus|no answer|ne sait|δεν ξέρω)", re.I)
DEMO_RE = re.compile(r"^(D\d|SD\d|DX\d|D\d+[a-z]|C\d+$)", re.I)
DEMO_TXT = re.compile(r"(how old are you|your age|gender|highest level of education|what is your current occupation|age when|marital|household|size of (the )?community|"
                      r"what is your nationality|how many people|which of the following best describes your|do you live in|region)", re.I)
FIELD = {"fl015ep": "2026-03", "fl574": "2026-01", "eb050ep": "2026-04", "eb049ep": "2025-11", "sp564": "2025-02", "fl563": "2025-04", "eb048ep": "2025-05",
         "eb047ep": "2025-01", "fl570": "2025-09"}


def subst(options):
    return [i for i, o in enumerate(options) if not DK_RE.search(o)]


def renorm(v, idx):
    x = np.array([v[i] if v[i] is not None else np.nan for i in idx], dtype=float)
    if np.isnan(x).any() or x.sum() <= 0:
        return None
    return x / x.sum()


def country_distances(root):
    """Κλειδωμένη απόσταση χωρών της Phase 5 (gbt-1): RMSE στις ορατές ελληνικές στήλες, χωρίς προσαρμογή σύνθεσης."""
    import gbt
    import gbt_eval as GE

    sv, stats, items = GE.load_all(root)
    countries, cols, Z, N = gbt.build_matrix(stats, items, use_adj=False)
    gr = countries.index("GR")
    mask = ~np.isnan(Z)
    vis = np.where(mask[gr])[0]
    d = {}
    for i, c in enumerate(countries):
        if c == "GR":
            continue
        o = vis[mask[i, vis]]
        if len(o) >= 10:
            d[c] = float(np.sqrt(np.mean((Z[i, o] - Z[gr, o]) ** 2)))
    return d


def volume_c_gaps(path_or_bytes, sheets_wanted):
    """Ανά ερώτηση: ridit διαφορές νέοι−ηλικιωμένοι, γυναίκες−άνδρες, υψηλή−χαμηλή μόρφωση από αρχείο Volume C μίας χώρας."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes, read_only=True, data_only=True)
    out = {}
    for sn in wb.sheetnames:
        if sn not in sheets_wanted:
            continue
        rows = list(wb[sn].iter_rows(values_only=True))
        h = next((i for i, r in enumerate(rows) if any(isinstance(c, str) and c.strip() in ("Gender", "Age") for c in r)), None)
        if h is None:
            continue
        grp, lab = rows[h], rows[h + 1]
        cur, cols = None, {}
        for j in range(len(lab)):
            if isinstance(grp[j], str) and grp[j].strip():
                cur = grp[j].strip()
            if isinstance(lab[j], str) and cur:
                cols[(cur, lab[j].strip())] = j
        def pick(group, pat):
            return [j for (g, l), j in cols.items() if g.lower().startswith(group) and re.search(pat, l, re.I)]
        G = {"female": pick("gender", r"^(woman|female)"), "male": pick("gender", r"^(man|male)"),
             "young": pick("age", r"^(15|18|16)\D*(24|29)|^25\D*(34|39)"), "old": pick("age", r"^(55|65)"),
             "edu_high": pick("educ", r"(high|20\+|tertiary|5-8)"), "edu_low": pick("educ", r"(low|15-|0-2|primary)")}
        opts, i = [], h + 2
        while i < len(rows) - 1:
            l = rows[i][1] if len(rows[i]) > 1 else None
            if isinstance(l, str) and l.startswith("Base"):
                i += 1
                continue
            if isinstance(l, str) and l.strip().lower().startswith("total"):
                break
            if isinstance(l, str) and l.strip():
                opts.append((l.strip(), rows[i + 1]))
                i += 2
                continue
            i += 1
        if len(opts) < 2:
            continue
        labels = [o for o, _ in opts]
        si = subst(labels)

        def dist(js):
            if not js:
                return None
            v = np.array([[float(r[j]) if isinstance(r[j], (int, float)) else 0.0 for j in js] for _, r in opts]).mean(axis=1)
            v = v[si]
            return v / v.sum() if v.sum() > 0 else None

        def ridit_gap(a, b):
            pa, pb = dist(G[a]), dist(G[b])
            if pa is None or pb is None:
                return None
            pm = (pa + pb) / 2
            rid = np.cumsum(pm) - pm / 2
            return float(pa @ rid - pb @ rid)

        out[sn] = {"young_old": ridit_gap("young", "old"), "female_male": ridit_gap("female", "male"), "edu_high_low": ridit_gap("edu_high", "edu_low")}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".")
    a = ap.parse_args()
    root = os.path.abspath(a.root)
    hd = os.path.join(root, ".synthetic-lab-data", "gpm", "holdout-h6")
    out_path = os.path.join(hd, "predictions.json")
    if os.path.exists(out_path):
        sys.exit("Οι προβλέψεις H6 υπάρχουν ήδη (μία εκτέλεση).")
    cfg = json.load(open(os.path.join(root, "config", "synthetic-lab", "persona", "gbt", "final-config.json")))
    dist = country_distances(root)
    import embed
    from gbt_eval import DOMAIN_PROTOTYPES
    from gbt_runtime import GlobalMatcher, extend_iqge
    from iqge_core import IQGE

    E = embed.Embedder(root)
    P = [E.encode(v) for v in DOMAIN_PROTOTYPES.values()]
    dnames = list(DOMAIN_PROTOTYPES)
    params = json.load(open(os.path.join(root, "config", "synthetic-lab", "persona", "iqge", "params.json")))
    mdir = os.path.join(root, ".synthetic-lab-data", "persona", "models", "pie-0.5")
    G = IQGE(root, mdir, params)
    GC_ = IQGE(root, mdir, params)
    extend_iqge(GC_, root)
    gm = GlobalMatcher(root, E)
    preds = {"createdAt": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "protocol": "PHASE6_PREREGISTRATION.md", "countryDistances": dist, "items": []}
    for f in sorted(glob.glob(os.path.join(hd, "visible", "*.json"))):
        sv = json.load(open(f))
        sid = sv["survey"]
        ref = FIELD[sid]
        for it in sv["items"]:
            opts = it["options"]
            si = subst(opts)
            if len(si) < 2:
                continue
            ctry = {c: renorm(v, si) for c, v in it["values"].items() if c not in AGG}
            ctry = {c: v for c, v in ctry.items() if v is not None}
            if len(ctry) < 5:
                continue
            raw_sum = np.mean([sum(x or 0 for x in v) for c, v in it["values"].items() if c not in AGG])
            sheet = it["id"].split("::")[1]
            demo = bool(DEMO_RE.match(sheet) or DEMO_TXT.search(it["question"]))
            multi = bool(raw_sum > 1.08)
            labels = [opts[i] for i in si]
            K = len(si)
            t = {"id": f"h6::{it['id']}", "text": re.sub(r"^\S+\s", "", it["question"]), "options": [{"code": str(i + 1), "label": l} for i, l in enumerate(labels)],
                 "type": "binary" if K == 2 else "ordinal", "wave": ref}
            v = E.encode([embed.item_text(t)])[0]
            dom = dnames[int(np.argmax([float((Pd @ v).max()) for Pd in P]))]
            pool = np.mean(list(ctry.values()), axis=0)
            nb = sorted([(dist[c], c) for c in ctry if c in dist])[:5]
            w = np.array([1 / max(x, 0.05) ** 2 for x, _ in nb])
            b5 = np.average([ctry[c] for _, c in nb], axis=0, weights=w)
            q80 = cfg["q80"].get(dom, cfg["q80"]["ALL"])
            rec = {"id": it["id"], "survey": sid, "question": it["question"][:300], "labels": labels, "substIdx": si, "K": K, "domain": dom, "demo": demo, "multi": multi,
                   "nCountries": len(ctry), "neighbors": [c for _, c in nb],
                   "UNIF": (np.ones(K) / K * 100).round(3).tolist(), "POOL": (pool * 100).round(3).tolist(), "B5": (b5 * 100).round(3).tolist(),
                   "B5_int80": [[max(0.0, x * 100 - q80), min(100.0, x * 100 + q80)] for x in b5], "disagreement": float(np.mean(np.abs(b5 - pool)) * 100)}
            # Κατηγορία A: ισοδύναμη ερώτηση με ΠΡΑΓΜΑΤΙΚΟ ελληνικό περιθώριο (κανόνας S1)
            if not multi and not demo and t["type"] != "nominal":
                m = gm.match(t)
                if m:
                    lp = gm.link_parts(m, labels)
                    intl = lp.get("international") or {}
                    if intl.get("basis") == "greek_observed":
                        rec["PERSIST"] = [round(lp["level"][l] * 100, 3) for l in labels]
                        rec["persistSource"] = intl.get("itemId")
            # Κατηγορία C: χωρίς τιμές άλλων χωρών
            if not multi and not demo:
                for name, g in (("IQGE", G), ("IQGE_C", GC_)):
                    r = g.interpret(t, ref)
                    lv = g.link(r, labels)["level"]
                    rec[name] = [round(lv[l] * 100, 3) for l in labels]
                    rec[name + "_tier"] = r["tier"]
            preds["items"].append(rec)
        print(sid, sum(1 for x in preds["items"] if x["survey"] == sid), flush=True)
    # H6.6: διαφορές υποομάδων από τα Volume C των ΑΛΛΩΝ χωρών
    preds["subgroup"] = {}
    for sid in ("fl015ep", "eb047ep"):
        pz = os.path.join(root, ".synthetic-lab-data", "raw", "phase6", "holdout", sid, "volumeC.zip")
        if not os.path.exists(pz):
            continue
        z = zipfile.ZipFile(pz)
        sheets = {x["id"].split("::")[1] for x in preds["items"] if x["survey"] == sid and not x["demo"] and not x["multi"]}
        per = {}
        for n in z.namelist():
            mm = re.search(r"_([A-Z]{2})\.xlsx$", n)
            if not mm or mm.group(1) in ("EL", "GR"):
                continue
            try:
                per[mm.group(1)] = volume_c_gaps(z.read(n), sheets)
            except Exception as ex:
                print("Volume C", sid, n, repr(ex)[:100])
        mean = {}
        for s in sheets:
            for k in ("young_old", "female_male", "edu_high_low"):
                vals = [per[c][s][k] for c in per if s in per[c] and per[c][s].get(k) is not None]
                if len(vals) >= 5:
                    mean.setdefault(s, {})[k] = float(np.mean(vals))
        preds["subgroup"][sid] = {"countries": sorted(per), "meanGap": mean}
    body = json.dumps(preds, sort_keys=True, ensure_ascii=False)
    preds["lockSha256"] = hashlib.sha256(body.encode()).hexdigest()
    json.dump(preds, open(out_path, "x"), ensure_ascii=False)
    print("κλείδωμα", preds["lockSha256"], "ερωτήσεις", len(preds["items"]))


if __name__ == "__main__":
    main()
