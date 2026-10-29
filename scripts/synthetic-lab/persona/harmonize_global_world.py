"""
Phase 5 GLOBAL BEHAVIOUR EVIDENCE LAKE — multi-country ("world") surveys, ALL countries.

    python3 -I scripts/synthetic-lab/persona/harmonize_global_world.py --root . [--only <surveyId> ...]

surveyIds: wrp2019 wrp2021 wrp2023 wrp2025 wgm2018 wgm2020 tisp covidistress lits4 qog2021 qog2024 maple_w12 maple_w34
           tealhelix populism2022

Output format: docs/synthetic-lab/GLOBAL_LAKE_SPEC.md. Writes ONLY under .synthetic-lab-data/global-behavior-evidence/v1/:
  resp/<sid>.csv.gz        rid, country (ISO2), wave (YYYY-MM), weight (mean 1 within country), age_band, gender, edu3, urban,
                           [pid: hashed panel link, MAPLE only], then "<sid>::<var>" = answer code as text (DK/refused/NA -> empty)
  locked/<sid>.GR.csv.gz   LOCKED final-holdout Greek rows (wgm2018 = F1, qog2021 = F2), same columns; absent from resp/
  items/<sid>.json         one object per item (same set for all countries)
  provenance/<sid>.json    source files with sha256, harmonising function, exclusions with reasons, notes
  provenance/_overlap.world.json   respondent-overlap check between files (ids / years only)
  surveys.world.json       survey entries (merged by surveyId on partial runs)

PREREGISTRATION (docs/synthetic-lab/PHASE5_PREREGISTRATION.md §1-§2):
  * Greek rows of wgm2018 (F1) and qog2021 (F2) are written only to locked/. For them NOTHING is computed except the row count:
    item inclusion (options, >=300 valid answers, >=2 distinct codes), undocumented-code checks and all printed counts use the
    non-locked rows only; option sets for these two surveys come from documentation (data dictionary / codebook / value labels),
    never from the data.
  * Political protection: items whose wording mentions vote / party / candidate / leader (automatic text filter, plus explicit
    MAPLE party/vote variables) carry excludedFromTransfer: true. They are kept, not transferred across countries.

Conventions reused from harmonize_p4_lake.py / harmonize_p4_holdouts.py / harmonize_p3.py / harmonize.py (imported, not modified):
  age bands 17-24 ... 65+ (rows aged <17 dropped where age exists), gender male/female, edu3 low (<=ISCED 2) / medium (ISCED 3-4) /
  high (ISCED 5+), DK/refusal/not-applicable -> empty, end-labelled numeric scales completed, multi-mention -> binary items,
  GDPR art. 9 special-category variables about the respondent (health, disability, ethnicity, sexual orientation, trade-union
  membership, discrimination grounds) are not items, producer-derived indices/recodes are not items.

DATA HYGIENE: never prints answer distributions, correlations or crosstabs - only row / item / valid-answer / country counts.
Downloads are read as data only (pandas / pyreadstat / openpyxl / python-docx / pdftotext); nothing inside them is executed.
Files are processed one at a time with usecols.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import unicodedata
import zipfile

import numpy as np
import pandas as pd
import pyreadstat

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import harmonize_p4_lake as L  # noqa: E402  (helpers + constants; module not modified)
import harmonize_p4_holdouts as H  # noqa: E402  (verified wording/codes; module not modified)

OUT_REL = os.path.join(".synthetic-lab-data", "global-behavior-evidence", "v1")
RAW4 = os.path.join(".synthetic-lab-data", "raw", "phase4")
RAWS = os.path.join(".synthetic-lab-data", "raw", "sprint", "microdata")
RAWP = os.path.join(".synthetic-lab-data", "raw", "political", "repos")
SCRIPT_REL = "scripts/synthetic-lab/persona/harmonize_global_world.py"
SURVEYS_FILE = "surveys.world.json"
MIN_VALID = 300
LOCKED = {"wgm2018": "F1", "qog2021": "F2"}
DEMO_OUT = ["rid", "country", "wave", "weight", "age_band", "gender", "edu3", "urban"]
INCLUSION_RULE = ("Every substantive attitude / behaviour / evaluation / knowledge variable asked in the common questionnaire, with "
                  "2-11 substantive options (after removing DK / refused / not applicable) and >= 300 valid answers over all "
                  "non-locked rows, with wording from the variable label, data dictionary, codebook or master questionnaire. "
                  "Excluded: demographics, IDs, admin/paradata, dates, weights, open text, producer-derived indices/recodes/scores, "
                  "coded open-ended answers, attention/manipulation checks, survey-experiment outcomes, country-specific question "
                  "variants, GDPR art. 9 special-category variables about the respondent, items whose option meanings are "
                  "undocumented.")

# ============================================================================== countries

NAME2ISO = {
    "afghanistan": "AF", "albania": "AL", "algeria": "DZ", "andorra": "AD", "angola": "AO", "antigua and barbuda": "AG",
    "argentina": "AR", "armenia": "AM", "australia": "AU", "austria": "AT", "azerbaijan": "AZ", "bahrain": "BH", "bangladesh": "BD",
    "barbados": "BB", "belarus": "BY", "belgium": "BE", "belize": "BZ", "benin": "BJ", "bhutan": "BT", "bolivia": "BO",
    "bosnia herzegovina": "BA", "bosnia and herzegovina": "BA", "botswana": "BW", "brazil": "BR", "brunei": "BN", "bulgaria": "BG",
    "burkina faso": "BF", "burundi": "BI", "cabo verde": "CV", "cambodia": "KH", "cameroon": "CM", "canada": "CA", "chad": "TD",
    "chile": "CL", "china": "CN", "colombia": "CO", "comoros": "KM", "congo brazzaville": "CG", "republic of congo": "CG",
    "congo republic of the": "CG", "congo kinshasa": "CD", "congo dr": "CD", "congo democratic republic of the": "CD",
    "costa rica": "CR", "croatia": "HR", "cuba": "CU", "cyprus": "CY", "czech republic": "CZ", "czechia": "CZ", "denmark": "DK",
    "djibouti": "DJ", "dominica": "DM", "dominican republic": "DO", "east timor timor leste": "TL", "ecuador": "EC", "egypt": "EG",
    "el salvador": "SV", "eritrea": "ER", "estonia": "EE", "eswatini": "SZ", "ethiopia": "ET", "fiji": "FJ", "finland": "FI",
    "france": "FR", "gabon": "GA", "gambia": "GM", "the gambia": "GM", "georgia": "GE", "germany": "DE", "ghana": "GH",
    "greece": "GR", "grenada": "GD", "guatemala": "GT", "guinea": "GN", "guinea bissau": "GW", "guyana": "GY", "haiti": "HT",
    "honduras": "HN", "hong kong": "HK", "hungary": "HU", "iceland": "IS", "india": "IN", "indonesia": "ID", "iran": "IR",
    "iraq": "IQ", "ireland": "IE", "israel": "IL", "italy": "IT", "ivory coast": "CI", "cote d ivoire": "CI", "c te d ivoire": "CI",
    "jamaica": "JM", "japan": "JP", "jordan": "JO", "kazakhstan": "KZ", "kenya": "KE", "kiribati": "KI", "korea north": "KP",
    "korea south": "KR", "south korea": "KR", "kosovo": "XK", "kuwait": "KW", "kyrgyzstan": "KG", "kyrgyz republic": "KG",
    "laos": "LA", "latvia": "LV", "lebanon": "LB", "lesotho": "LS", "liberia": "LR", "libya": "LY", "liechtenstein": "LI",
    "lithuania": "LT", "luxembourg": "LU", "macedonia": "MK", "north macedonia": "MK", "madagascar": "MG", "malawi": "MW",
    "malaysia": "MY", "maldives": "MV", "mali": "ML", "malta": "MT", "mauritania": "MR", "mauritius": "MU", "mexico": "MX",
    "micronesia federated states of": "FM", "moldova": "MD", "monaco": "MC", "mongolia": "MN", "montenegro": "ME", "morocco": "MA",
    "mozambique": "MZ", "myanmar": "MM", "myanmar burma": "MM", "namibia": "NA", "nauru": "NR", "nepal": "NP", "netherlands": "NL",
    "new zealand": "NZ", "nicaragua": "NI", "niger": "NE", "nigeria": "NG", "norway": "NO", "oman": "OM", "pakistan": "PK",
    "palestine": "PS", "palestinian territories": "PS", "west bank and gaza": "PS", "panama": "PA", "paraguay": "PY", "peru": "PE",
    "philippines": "PH", "poland": "PL", "portugal": "PT", "qatar": "QA", "romania": "RO", "russia": "RU", "rwanda": "RW",
    "saint lucia": "LC", "saint vincent and the grenadines": "VC", "san marino": "SM", "sao tome and principe": "ST",
    "saudi arabia": "SA", "senegal": "SN", "serbia": "RS", "seychelles": "SC", "sierra leone": "SL", "singapore": "SG",
    "slovakia": "SK", "slovak republic": "SK", "slovenia": "SI", "solomon islands": "SB", "somalia": "SO", "south africa": "ZA",
    "spain": "ES", "sri lanka": "LK", "sudan": "SD", "sudan south": "SS", "suriname": "SR", "sweden": "SE", "switzerland": "CH",
    "taiwan": "TW", "tajikistan": "TJ", "tanzania": "TZ", "thailand": "TH", "the bahamas": "BS", "togo": "TG",
    "trinidad and tobago": "TT", "tunisia": "TN", "turkey": "TR", "turkiye": "TR", "turkmenistan": "TM", "uganda": "UG",
    "ukraine": "UA", "united arab emirates": "AE", "united kingdom": "GB", "united states": "US", "uruguay": "UY",
    "uzbekistan": "UZ", "venezuela": "VE", "vietnam": "VN", "yemen": "YE", "zambia": "ZM", "zimbabwe": "ZW",
}
ISO3 = {"ALB": "AL", "ARG": "AR", "AUS": "AU", "AUT": "AT", "BEL": "BE", "BGD": "BD", "BGR": "BG", "BOL": "BO", "BRA": "BR",
        "BWA": "BW", "CAN": "CA", "CHE": "CH", "CHL": "CL", "CHN": "CN", "CIV": "CI", "CMR": "CM", "COD": "CD", "COL": "CO",
        "CRI": "CR", "CYP": "CY", "CZE": "CZ", "DEU": "DE", "DNK": "DK", "EGY": "EG", "ESP": "ES", "ETH": "ET", "FIN": "FI",
        "FRA": "FR", "GBR": "GB", "GEO": "GE", "GHA": "GH", "GRC": "GR", "HKG": "HK", "HUN": "HU", "IDN": "ID", "IND": "IN",
        "IRL": "IE", "ISR": "IL", "ITA": "IT", "JPN": "JP", "KAZ": "KZ", "KEN": "KE", "KOR": "KR", "MAR": "MA", "MEX": "MX",
        "MYS": "MY", "NGA": "NG", "NIC": "NI", "NLD": "NL", "NOR": "NO", "NZL": "NZ", "PER": "PE", "PHL": "PH", "POL": "PL",
        "PRT": "PT", "ROU": "RO", "RUS": "RU", "SRB": "RS", "SVK": "SK", "SVN": "SI", "SWE": "SE", "TUR": "TR", "TWN": "TW",
        "UGA": "UG", "UKR": "UA", "URY": "UY", "USA": "US", "ZAF": "ZA"}
NUTS2ISO = {"EL": "GR", "UK": "GB"}


def name_key(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower().replace("&", " and ")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def iso2(name):
    return NAME2ISO.get(name_key(name))


# ============================================================================== generic helpers

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def code_str(v):
    return L.code_str(v)


def ym_from_stata_ms(v):
    """Stata %tc (ms since 1960-01-01) -> 'YYYY-MM'."""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f != f:
        return None
    d = dt.datetime(1960, 1, 1) + dt.timedelta(milliseconds=f)
    return f"{d.year:04d}-{d.month:02d}"


def ym_series(s):
    d = pd.to_datetime(pd.Series(s), errors="coerce")
    return d.dt.strftime("%Y-%m").where(d.notna(), None)


def country_waves(country, ym, default):
    """{ISO2: earliest 'YYYY-MM' holding >=1% of the country's dated rows}; default where nothing is dated."""
    df = pd.DataFrame({"c": pd.Series(country).values, "m": pd.Series(ym).values})
    out = {}
    for c, g in df.groupby("c"):
        m = L.earliest_month(g.m)
        out[c] = m or default
    return out


def age_band(a):
    return L.age_band(a)


EXTRA_NONSUB = re.compile(r"difficult to say|^none of (these|the above)|cannot provide|haven.?t heard|don.?t read newspapers|"
                          r"never attended|not eligible|does not apply|no opinion", re.I)


def vl_mark(vl):
    """Value labels with additional non-substantive labels (LiTS 'Difficult to say', 'None of these', ...) prefixed so that
    harmonize_p4_lake.NONSUB recognises them."""
    out = {}
    for k, v in (vl or {}).items():
        lab = str(v)
        if EXTRA_NONSUB.search(L.norm(lab)) and not L.NONSUB.search(L.norm(lab)):
            lab = f"(n/a) {lab}"
        out[k] = lab
    return out


def skip_reason(var, label, sensitive=True, demographic=True):
    """harmonize_p4_lake.generic_skip, with the GDPR art. 9 filter optional (LiTS uses an explicit list instead)."""
    if sensitive:
        return L.generic_skip(var, label)
    core = L.QID_PREFIX.sub("", label)
    if L.ADMIN_LABEL.search(label) or L.ADMIN_LABEL.search(core):
        return "admin / ID / weight / date / paradata"
    if L.DERIVED_LABEL.search(label):
        return "producer-derived index / recode / net / total"
    if demographic and (L.DEMO_LABEL.search(label) or L.DEMO_LABEL.search(core)):
        return "demographic"
    return None


# political protection (PHASE5 prereg §2): automatic text filter
LEADER_OK = re.compile(r"(religious|community|traditional|business|spiritual|police/community|famous)\s+leaders?", re.I)
POLITICAL_RX = re.compile(r"\b(vot(e|es|ed|er|ers|ing)|ballot|part(y|ies)|partisan\w*|candidates?|leaders?|presiden\w*|"
                          r"prime minister)\b", re.I)
POLITICAL_EL = re.compile(r"(^|[^α-ω])κομμα|ψηφ(?!ιακ)|υποψηφι|αρχηγ|πρωθυπουργ")  # on L.norm text (no accents); not αποκομμένος / ψηφιακός
POLITICAL_REASON = ("political protection (PHASE5_PREREGISTRATION §2): wording mentions vote / party / candidate / leader; kept, "
                    "not transferred across countries")


def is_political(text):
    t = LEADER_OK.sub(" ", str(text or ""))
    return bool(POLITICAL_RX.search(t) or POLITICAL_EL.search(L.norm(t)))


def five(lo, hi):
    return [("1", lo), ("2", "2"), ("3", "3"), ("4", "4"), ("5", hi)]


def map_codes(series, options, nonsub=()):
    """numeric codes -> str codes; nonsub codes -> None; undocumented codes are kept (finalize nulls and counts them)."""
    codes = pd.to_numeric(series, errors="coerce").map(code_str).astype(object)
    ns = {str(c) for c in nonsub}
    codes = codes.where(codes.notna() & ~codes.isin(ns), None)
    return codes


def add_item(ctx, var, codes, text, options, *, language="en", qtype=None, label=None, text_source="", stem=None, nonsub=(),
             split="all", political=False, multi=False):
    ctx.add(var, codes, text, options, nonsub=nonsub, qtype=qtype, label=label, text_source=text_source, stem=stem,
            multi=multi, language=language)
    ctx.items[-1]["split"] = split
    ctx.items[-1]["political"] = political


def tag_last(ctx, **kw):
    if ctx.items:
        ctx.items[-1].update(kw)


def new_demo(n):
    return pd.DataFrame({"country": [None] * n, "wave": [None] * n, "weight": np.nan, "age_band": [None] * n, "gender": [None] * n,
                         "edu3": [None] * n, "urban": [None] * n})


def gallup_edu(s):
    return pd.to_numeric(s, errors="coerce").map({1.0: "low", 2.0: "medium", 3.0: "high"})


def mf(s, male, female):
    return pd.to_numeric(s, errors="coerce").map({float(male): "male", float(female): "female"})


# ============================================================================== finalisation / output

def out_dir(root, *parts):
    p = os.path.join(root, OUT_REL, *parts)
    os.makedirs(os.path.dirname(p) if os.path.splitext(p)[1] else p, exist_ok=True)
    return p


def finalize(root, sid, ctx, demo, meta, sources, *, pid=None, func=None):
    """Thresholds, political flags, weight normalisation, writes resp/locked/items/provenance; returns surveys.json entry."""
    demo = demo.reset_index(drop=True)
    n = len(demo)
    if demo.country.isna().any():
        raise ValueError(f"{sid}: rows without country reached finalize")
    locked = sid in LOCKED
    lmask = (demo.country == "GR") if locked else pd.Series(False, index=demo.index)
    omask = ~lmask
    # weights: source national weight, normalised to mean 1 within country (1 where no weight)
    w = pd.to_numeric(demo.weight, errors="coerce").astype(float)
    w = w.where(w > 0)
    wn = pd.Series(1.0, index=demo.index)
    for c, idx in demo.groupby("country").groups.items():
        s = w.loc[idx]
        if s.notna().any():
            s = s.fillna(s.mean())
            wn.loc[idx] = s / s.mean()
    demo["weight"] = wn.round(6)
    demo.insert(0, "rid", [f"{sid}-{i + 1}" for i in range(n)])
    cols, items_out, kept = {}, [], []
    for it in ctx.items:
        codes = pd.Series(it["codes"]).reset_index(drop=True).astype(object)
        if len(codes) != n:
            raise ValueError(f"{sid}::{it['var']}: length mismatch")
        codes = codes.where(codes.notna(), None)
        valid = {c for c, _ in it["options"]}
        bad = codes.notna() & ~codes.isin(valid)
        n_bad_open = int((bad & omask).sum())
        codes = codes.where(~bad, None)  # undocumented codes -> empty (locked rows: silently, never counted)
        k = len(it["options"])
        if not 2 <= k <= 11:
            ctx.exclude(it["var"], f"{k} substantive options (rule: 2-11)")
            continue
        oc = codes[omask]
        n_valid = int(oc.notna().sum())
        if n_valid < MIN_VALID:
            ctx.exclude(it["var"], f"{n_valid} valid answers{' (non-locked rows)' if locked else ''} (<{MIN_VALID})")
            continue
        if oc.dropna().nunique() < 2:
            ctx.exclude(it["var"], "fewer than 2 distinct answer codes observed (constant / routing artefact)")
            continue
        if n_bad_open:
            ctx.notes.append(f"{it['var']}: {n_bad_open} answers with codes outside the documented options set to empty")
        n_cty = int(demo.country[omask & codes.notna()].nunique())
        if n_cty == 1 and demo.country[omask].nunique() >= 3:
            ctx.exclude(it["var"], "answered in one country only (country-specific question, not part of the common questionnaire)")
            continue
        it["codes"], it["nValid"] = codes, n_valid
        it["nCountries"] = n_cty
        kept.append(it)
    if ctx.empties:
        ctx.excluded.append({"variable": f"{len(ctx.empties)} variables", "reason": "no answers at all", "variables": ctx.empties})
    flags = [f and not re.search(r"[?;\u037e]\s*[»”\"]?\s*$", it["text"]) for f, it in zip(L.battery_flags(kept), kept)]
    for it, flag in zip(kept, flags):
        iid = f"{sid}::{it['var']}"
        cols[iid] = it["codes"]
        pol = bool(it.get("political")) or is_political(it["text"]) or is_political(it.get("label") or "")
        entry = {"id": iid, "var": it["var"], "text": it["text"], "textLanguage": it["language"],
                 "options": [{"code": c, "label": l} for c, l in it["options"]], "type": it["type"],
                 "stemMissing": bool(flag), "split": it.get("split", "all"), "excludedFromTransfer": pol,
                 "excludedReason": (it.get("politicalReason") or POLITICAL_REASON) if pol else None,
                 "label": it["label"], "textSource": it["textSource"], "nonsubCodes": it["nonsub"], "nValid": it["nValid"],
                 "nCountries": it["nCountries"]}
        if it.get("multi"):
            entry["multiMention"] = True
        if it.get("stem"):
            entry["stemAdded"] = True
        if it.get("note"):
            entry["note"] = it["note"]
        items_out.append(entry)
    base = demo[DEMO_OUT].copy()
    if pid is not None:
        base["pid"] = pd.Series(pid).reset_index(drop=True).values
    frame = pd.concat([base, pd.DataFrame(cols)], axis=1)
    resp = frame[omask.values]
    if locked:
        assert not (resp.country == "GR").any(), f"{sid}: locked Greek rows leaked into resp"
    resp.to_csv(out_dir(root, "resp", f"{sid}.csv.gz"), index=False, compression="gzip")
    lpath = os.path.join(root, OUT_REL, "locked", f"{sid}.GR.csv.gz")
    n_locked = int(lmask.sum())
    if locked:
        frame[lmask.values].to_csv(out_dir(root, "locked", f"{sid}.GR.csv.gz"), index=False, compression="gzip")
    elif os.path.exists(lpath):
        os.remove(lpath)
    with open(out_dir(root, "items", f"{sid}.json"), "w", encoding="utf-8") as fh:
        json.dump(items_out, fh, ensure_ascii=False, indent=1)
    # surveys.json entry
    demo["wave"] = demo.wave.map(lambda v: v if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}", v) else None)
    cw = demo.groupby("country").wave.agg(lambda s: L.earliest_month(s.dropna()) or None).to_dict()
    cw = {k: (v if isinstance(v, str) else None) for k, v in cw.items()}
    cn = demo.groupby("country").size().to_dict()
    countries = {c: {"n": int(cn[c]), "wave": cw.get(c)} for c in sorted(cn)}
    waves = sorted(v["wave"] for v in countries.values() if v["wave"])
    greek_role = ("final_locked" if locked else "dev") if "GR" in countries else "absent"
    entry = {"surveyId": sid, "title": meta["title"], "family": meta["family"], "source": meta["source"],
             "doi_or_url": meta["doi_or_url"], "licence": meta["licence"], "commercialUse": meta["commercialUse"],
             "aiTraining": meta["aiTraining"], "wave": waves[0] if waves else None, "waveApprox": meta["waveApprox"],
             "mode": meta["mode"], "sampling": meta["sampling"], "weightVar": meta["weightVar"], "countries": countries,
             "nItems": len(items_out), "greekRole": greek_role, "notes": list(meta.get("notes", [])) + ctx.notes}
    if locked:
        entry["lockedHoldout"] = LOCKED[sid]
        entry["notes"].append(f"Greek rows are locked final holdout {LOCKED[sid]}: written only to locked/{sid}.GR.csv.gz "
                              f"(n={n_locked}); nothing but the row count is computed for them.")
    if "NA" in countries:
        entry["notes"].append("Namibia's ISO2 code is 'NA': read the CSV with keep_default_na=False (or na_filter=False) so that "
                              "it is not parsed as missing.")
    entry["nRows"] = {"resp": int(omask.sum()), "locked": n_locked}
    entry["textLanguage"] = sorted({it["textLanguage"] for it in items_out})
    prov = {"surveyId": sid, "generatedAt": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "harmonizer": {"script": SCRIPT_REL, "function": func, "command": f"python3 -I {SCRIPT_REL} --root . --only {sid}"},
            "sources": sources, "inclusionRule": INCLUSION_RULE, "minValidAnswers": MIN_VALID,
            "rows": entry["nRows"], "excluded": ctx.excluded, "notes": entry["notes"],
            "politicalFilter": POLITICAL_RX.pattern + " | Greek: " + POLITICAL_EL.pattern + " (religious/community/traditional/"
                               "business/famous leaders ignored)"}
    for k in ("weightNote", "overlap"):
        if k in meta:
            prov[k] = meta[k]
    if meta.get("waveApproxCountries"):
        entry["waveApproxCountries"] = meta["waveApproxCountries"]
    with open(out_dir(root, "provenance", f"{sid}.json"), "w", encoding="utf-8") as fh:
        json.dump(prov, fh, ensure_ascii=False, indent=1)
    n_pol = sum(1 for it in items_out if it["excludedFromTransfer"])
    print(f"{sid}: resp rows={int(omask.sum())} countries={len(countries)} items={len(items_out)} (excludedFromTransfer={n_pol}) "
          f"excluded={len(ctx.excluded)}" + (f" | LOCKED GR rows={n_locked}" if locked else ""))
    return entry


def src(root, path):
    ap = path if os.path.isabs(path) else os.path.join(root, path)
    return {"path": os.path.relpath(ap, root), "sha256": sha256(ap), "bytes": os.path.getsize(ap)}


def man_phase4(root, slug):
    return L.load_manifest(root).get(slug, {})


def man_sprint(root, slug):
    return H.manifest_entry(root, slug)


# ============================================================================== World Risk Poll

WRP_EXTRA_SKIP = {"IncomeFeelings": "feelings about household income (socio-economic background; treated as demographic as in WGM)",
                  "WP1223": "demographic (marital status)", "EMP8B": "employment background (occupation)",
                  "WP23336": "employment background (last time employed)", "WP23340": "employment background (job sector, coded)"}


def h_wrp(root, year):
    sid = f"wrp{year}"
    slug = "lrf_world_risk_poll_2019_2025"
    man = man_phase4(root, slug)
    path = os.path.join(root, RAW4, slug, "extracted", f"WRP_{year}.sav")
    _, m = pyreadstat.read_sav(path, metadataonly=True)
    labs, vls = m.column_names_to_labels or {}, m.variable_value_labels or {}
    ctx = L.Ctx(sid, slug)
    cand = []
    for var in m.column_names:
        lab = L.clean(labs.get(var) or "")
        if L.WRP_SKIP.match(var) and var not in ("FIELD_DATE",):
            ctx.exclude(var, "harmonised-block demographic / admin / weight / derived score (skip list)")
            continue
        if var in WRP_EXTRA_SKIP:
            ctx.exclude(var, WRP_EXTRA_SKIP[var])
            continue
        if var == "FIELD_DATE":
            continue
        if var.endswith("_ALL"):
            ctx.exclude(var, "producer-combined version of country wording variants (the common base variable is kept)")
            continue
        r = L.generic_skip(var, lab)
        if r:
            ctx.exclude(var, r)
            continue
        if re.search(r"coded response|\(coded\)", lab, re.I):
            ctx.exclude(var, "coded open-ended response")
            continue
        cand.append(var)
    usecols = ["COUNTRY_ISO2", "Age", "Gender", "Education", "Urbanicity2", "WGT", "WPID_RANDOM"] + \
              (["FIELD_DATE"] if "FIELD_DATE" in m.column_names else []) + cand
    df, _ = pyreadstat.read_sav(path, usecols=usecols)
    n0 = len(df)
    df["COUNTRY_ISO2"] = df.COUNTRY_ISO2.astype(str).str.strip()
    nocty = ~df.COUNTRY_ISO2.str.fullmatch(r"[A-Z]{2}")
    age = pd.to_numeric(df.Age, errors="coerce")
    drop = nocty | (age < 17)
    df = df[~drop].reset_index(drop=True)
    age = pd.to_numeric(df.Age, errors="coerce")
    demo = new_demo(len(df))
    demo["country"] = df.COUNTRY_ISO2.values
    demo["weight"] = df.WGT.values
    demo["age_band"] = age.map(lambda a: age_band(a) if a == a and a < 100 else None)
    demo["gender"] = mf(df.Gender, 1, 2)
    demo["edu3"] = gallup_edu(df.Education)
    demo["urban"] = pd.to_numeric(df.Urbanicity2, errors="coerce").map({1.0: "rural", 2.0: "urban"})
    if "FIELD_DATE" in df:
        cw = country_waves(demo.country, ym_series(df.FIELD_DATE), f"{year}-07")
        wave_basis = "FIELD_DATE ('Study Completion Date') earliest month per country (>=1% of cases) - completion, not start date"
    else:
        cw = {c: f"{year}-07" for c in demo.country.unique()}
        wave_basis = "no date variable in the file: mid-year of the fieldwork year used for every country"
    demo["wave"] = demo.country.map(cw)
    ncty_total = demo.country.nunique()
    for var in cand:
        lab = L.clean(labs.get(var) or "")
        s = df[var]
        nc = int(demo.country[s.notna()].nunique())
        if 0 < nc <= 3 < ncty_total:
            ctx.exclude(var, f"country-specific wording variant (asked in {nc} countries); not part of the common questionnaire")
            continue
        L.add_labelled(ctx, var, s, lab, vl_mark(vls.get(var)), text=lab,
                       text_source="SPSS variable label (LRF/Gallup short wording; full questionnaire not in the deposit)")
        tag_last(ctx, split="all")
    meta = {"title": f"Lloyd's Register Foundation World Risk Poll {year}", "family": "wrp",
            "source": f"{man.get('producer')}; harmonised single-wave file WRP_{year}.sav (LRF build 2026-08-31)",
            "doi_or_url": man.get("url"), "licence": man.get("licence"), "commercialUse": "unclear", "aiTraining": "unclear",
            "waveApprox": True, "mode": "Gallup World Poll: face-to-face (CAPI/PAPI) or telephone depending on country (not in file)",
            "sampling": "Gallup World Poll national probability samples of residents aged 15+ (~1,000 per country)",
            "weightVar": "WGT (within-country survey weight; PROJWT is for pooled/projection analyses and is not used)",
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} (aged 15-16 or no ISO2 country).", f"wave: {wave_basis}.",
                      "edu3 from Gallup Education (primary or less = low; secondary 9-15 years = medium, includes some lower-"
                      "secondary completers; tertiary = high). urban from Urbanicity2 (rural/small town = rural; large city/"
                      "suburb = urban).", "Item text = LRF/Gallup variable label (short wording); English master wording.",
                      "Licence: free use with attribution (LRF data page); no explicit commercial / AI-training clause.",
                      "XK (Kosovo) is a user-assigned code, not an official ISO 3166-1 alpha-2 code."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path)], func=f"h_wrp(year={year})")


# ============================================================================== Wellcome Global Monitor 2018 (LOCKED GR = F1)

def parse_wgm_codes(s):
    s = str(s or "").replace("\xa0", " ")
    out = []
    for mm in re.finditer(r"(\d+)\s*[=\-]\s*(.+?)(?=,?\s+\d+\s*[=\-]\s*\S|,\s*\(|$)", s):
        out.append((mm.group(1), mm.group(2).strip().rstrip(",").strip()))
    return out


def wgm2018_methodology_months(path):
    """Table II.B of the 2018 methodology: '<d>[ <Mon>]–<d>[ <Mon>,]' on the line above '2018 <Country>', '[<Mon>, ]2018' below."""
    try:
        txt = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True, timeout=120).stdout
    except Exception:  # noqa: BLE001
        return {}
    lines = txt.splitlines()
    mon = {m[:3]: i for i, m in enumerate(["january", "february", "march", "april", "may", "june", "july", "august", "september",
                                            "october", "november", "december"], 1)}
    out = {}
    for i, ln in enumerate(lines):
        if i == 0 or i + 1 >= len(lines):
            continue
        above, below = lines[i - 1][:40], lines[i + 1][:40]
        mm = re.match(r"^\s*2018\s+([A-Z][A-Za-z .'’()-]+?)\s{2,}[\d,]+\s", ln)
        if not mm and re.match(r"^\s*2018\s+[\d,]+\s", ln):  # country name split over the lines above / below
            ma, mb = re.match(r"^\s*([A-Z][A-Za-z]*(?: [a-z]+)?)\s+(\d.*)$", above), re.match(r"^\s*([A-Z][A-Za-z]+)\b(.*)$", below)
            if ma and mb:
                mm = re.match(r"(.*)", f"{ma.group(1)} {mb.group(1)}")
                above, below = ma.group(2), mb.group(2)
        if not mm:
            continue
        dm = re.search(r"(\d{1,2})\s*([A-Z][a-z]{2})?\s*[–-]\s*\d{1,2}\s*([A-Z][a-z]{2})?", above)
        if not dm:
            continue
        cand = [dm.group(2), dm.group(3)] + re.findall(r"\b([A-Z][a-z]{2}),", below)
        cand = [c for c in cand if c and c.lower() in mon]
        if not cand:
            continue
        ym = re.search(r"(20\d\d)", below) or re.search(r"(20\d\d)", above)
        y = int(ym.group(1)) if ym else 2018
        sm = mon[cand[0].lower()]
        em = mon[cand[-1].lower()]
        if sm > em:
            y -= 1
        c = iso2(mm.group(1).strip()) or {"swaziland": "SZ"}.get(name_key(mm.group(1)))
        if c:
            out[c] = f"{y:04d}-{sm:02d}"
    return out


def h_wgm2018(root):
    sid = "wgm2018"
    base = os.path.join(root, RAWS, "wellcome_global_monitor_2018")
    path = os.path.join(base, "wgm2018-dataset-crosstabs-all-countries.xlsx")
    man = man_sprint(root, "wellcome_global_monitor_2018")
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    dd = {}
    for row in wb["Data dictionary"].iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            dd[str(row[0]).strip()] = (L.clean(row[1]), str(row[2] or ""), L.clean(row[3] if len(row) > 3 else ""))
    wb.close()
    wp5_names = dict((int(a), b.strip()) for a, b in re.findall(r"(\d+)=([^,]+)", dd["WP5"][1]))
    ctx = L.Ctx(sid, "wellcome_global_monitor_2018")
    item_vars = [v for v in dd if re.match(r"^Q\d+[A-Z]?$", v)]
    skip = {"Q27": "demographic (has children)"}
    for v, r in [("D1", "religion (recoded; GDPR art. 9 / demographic)"), ("WGM_Index", "producer-derived index"),
                 ("WGM_Indexr", "producer-derived index (recoded)"), ("ViewOfScience", "producer-derived typology of Q16/Q17"),
                 ("Age, AgeCategories, Gender, Education, Urban_Rural, Household_Income, Subjective_Income, EMP_2010",
                  "demographics / producer recodes"), ("WP5, wgt, PROJWT, FIELD_DATE, YEAR_CALENDAR, Regions_Report, WBI",
                                                       "country / weights / admin / country-level variables")]:
        ctx.exclude(v, r)
    usecols = ["WP5", "wgt", "Age", "Gender", "Education", "Urban_Rural"] + [v for v in item_vars if v not in skip]
    df = pd.read_excel(path, sheet_name="Full dataset", usecols=usecols, engine="openpyxl")
    n0 = len(df)
    cty = pd.to_numeric(df.WP5, errors="coerce").map(lambda x: iso2(wp5_names.get(int(x))) if x == x else None)
    age = pd.to_numeric(df.Age, errors="coerce").where(lambda a: a < 100)
    n_nocty = int(cty.isna().sum())
    drop = cty.isna() | (age < 17)
    df, cty, age = df[~drop].reset_index(drop=True), cty[~drop].reset_index(drop=True), age[~drop].reset_index(drop=True)
    demo = new_demo(len(df))
    demo["country"] = cty.values
    demo["weight"] = df.wgt.values
    demo["age_band"] = age.map(age_band)
    demo["gender"] = mf(df.Gender, 1, 2)
    demo["edu3"] = gallup_edu(df.Education)
    demo["urban"] = pd.to_numeric(df.Urban_Rural, errors="coerce").map({1.0: "rural", 2.0: "urban"})
    meth = wgm2018_methodology_months(os.path.join(base, "wgm2018-methodology.pdf"))
    cw = {c: meth.get(c) or "2018-07" for c in demo.country.unique()}
    demo["wave"] = demo.country.map(cw)
    n_from_meth = sum(1 for c in cw if c in meth)
    approx_c = sorted(c for c in cw if c not in meth)
    for v in item_vars:
        if v in skip:
            ctx.exclude(v, skip[v])
            continue
        text, codes_s, note = dd[v]
        parsed = parse_wgm_codes(codes_s)
        opts = [(c, l) for c, l in parsed if not (L.NONSUB.search(L.norm(l)) or EXTRA_NONSUB.search(L.norm(l)))]
        nons = [c for c, l in parsed if (L.NONSUB.search(L.norm(l)) or EXTRA_NONSUB.search(L.norm(l)))]
        if v in ("Q24", "Q25", "Q26"):
            nons = sorted(set(nons) | {"99"})
        add_item(ctx, v, map_codes(df[v], opts, nons), text, opts, nonsub=nons,
                 text_source="WGM 2018 data dictionary sheet ('Variable Label/Survey Question'; English master questionnaire)")
        if note:
            tag_last(ctx, note=note)
    meta = {"title": "Wellcome Global Monitor 2018", "family": "wgm",
            "source": "Wellcome Trust / Gallup World Poll 2018 add-on; 'Full dataset' sheet of wgm2018-dataset-crosstabs-all-countries.xlsx",
            "doi_or_url": man.get("doi_or_url"),
            "licence": ("Caller states CC BY 4.0, NOT verified: no licence text in the local files (data dictionary, methodology, "
                        "questionnaire). Manifest: " + str(man.get("licence"))),
            "commercialUse": "unclear", "aiTraining": "unclear", "waveApprox": False,
            "mode": "Gallup World Poll 2018: face-to-face CAPI/PAPI or telephone by country (methodology Table II.B)",
            "sampling": "Gallup World Poll national probability samples of residents aged 15+ (~1,000 per country)",
            "weightVar": "wgt (national weight for analysis at country level; PROJWT not used)",
            "waveApproxCountries": approx_c,
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} ({n_nocty} Northern Cyprus / unmapped WP5 rows - not an ISO "
                      "3166-1 country - and respondents aged 15-16).",
                      f"wave = fieldwork start month from methodology Table II.B for {n_from_meth} countries; 2018-07 (mid-year "
                      f"approximation) for {approx_c}. FIELD_DATE is not used: its values look day/month-swapped (e.g. Greek cases "
                      "2018-01-07 while the methodology gives 18 Jun-17 Jul 2018).",
                      "Options from the data dictionary codes; (DK)/(Refused)/(Never attended this type of school) are empty.",
                      "edu3/urban from Gallup Education / Urban_Rural recodes. Country from WP5 via the dictionary's country names."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path), src(root, os.path.join(base, "wgm2018-methodology.pdf")),
                                                 src(root, os.path.join(base, "wgm2018-questionnaire.pdf"))], func="h_wgm2018")


# ============================================================================== Wellcome Global Monitor 2020

WGM20_ART9 = {"MH7A", "MH7C"} | {f"MH8{c}" for c in "ABCDEFGH"} | {f"MH9{c}" for c in "ABCDEFGH"}


def h_wgm2020(root):
    sid = "wgm2020"
    path = os.path.join(root, RAWS, "wellcome_global_monitor_2020", "wgm-full-wave2-public-file.sav")
    man = man_sprint(root, "wellcome_global_monitor_2020")
    _, m = pyreadstat.read_sav(path, metadataonly=True)
    vls = m.variable_value_labels or {}
    ctx = L.Ctx(sid, "wellcome_global_monitor_2020")
    items = [v for v in H.WGM_TEXT if v not in WGM20_ART9]
    for v in sorted(WGM20_ART9):
        ctx.exclude(v, "respondent's own anxiety/depression episodes and coping/treatment (GDPR art. 9 health data)")
    for v, r in [("W3", "standardized recode of education where respondent last learned about science (producer recode)"),
                 ("MH7B", "open-ended age (continuous)"), ("MH7B_2, age_mh", "age at first anxiety/depression (health; derived)"),
                 ("Subjective_Income, Household_Income, EMP_2010, Education, Gender, Age, age_var1-3", "demographics / recodes"),
                 ("COUNTRYNEW, WPID_RANDOM, WGT, PROJWT, FIELD_DATE, YEAR_WAVE, Global11Regions, wbi", "ID / weight / admin")]:
        ctx.exclude(v, r)
    df, _ = pyreadstat.read_sav(path, usecols=["COUNTRYNEW", "WGT", "FIELD_DATE", "Age", "Gender", "Education"] + items)
    n0 = len(df)
    cty = df.COUNTRYNEW.map(iso2)
    age = pd.to_numeric(df.Age, errors="coerce").where(lambda a: a < 100)
    drop = cty.isna() | (age < 17)
    unm = sorted(set(df.COUNTRYNEW[cty.isna()]))
    df, cty, age = df[~drop].reset_index(drop=True), cty[~drop].reset_index(drop=True), age[~drop].reset_index(drop=True)
    demo = new_demo(len(df))
    demo["country"] = cty.values
    demo["weight"] = df.WGT.values
    demo["age_band"] = age.map(age_band)
    demo["gender"] = mf(df.Gender, 1, 2)
    demo["edu3"] = gallup_edu(df.Education)
    demo["wave"] = demo.country.map(country_waves(demo.country, ym_series(df.FIELD_DATE), "2020-09"))
    src_txt = ("Wellcome Global Monitor 2020 full questionnaire (English master; wording verified in harmonize_p4_holdouts.WGM_TEXT, "
               "'Greece' replaced by '(country)') + SPSS value labels")
    for v in items:
        text, qtype = H.WGM_TEXT[v]
        text = text.replace("Greece", "(country)")
        L.add_labelled(ctx, v, df[v], L.clean(m.column_names_to_labels.get(v) or v), vl_mark(vls.get(v)), text=text,
                       text_source=src_txt)
    meta = {"title": "Wellcome Global Monitor 2020 (Covid-19 wave)", "family": "wgm",
            "source": "Wellcome Trust / Gallup World Poll 2020; wgm-full-wave2-public-file.sav", "doi_or_url": man.get("doi_or_url"),
            "licence": ("Caller states CC BY 4.0, NOT verified: no licence text in the local data dictionary docx. Manifest: "
                        + str(man.get("licence"))),
            "commercialUse": "unclear", "aiTraining": "unclear", "waveApprox": True,
            "mode": "Gallup World Poll 2020: mostly telephone (Covid-19 period), face-to-face in some countries",
            "sampling": "Gallup World Poll national probability samples of residents aged 15+",
            "weightVar": "WGT (national weight for country-level analysis; PROJWT not used)",
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} (aged 15-16 or unmapped country {unm}).",
                      "wave = earliest FIELD_DATE ('Study Completion Date') month per country (>=1% of cases); approximation.",
                      "'(Does not apply/No job)' in WP21758-61 and '(Never heard of it)' handling: does-not-apply is empty "
                      "(spec: NA -> empty; harmonize_p4_holdouts kept it as substantive); '(Never heard of it)' kept.",
                      "No urban/rural variable in the public file."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path), src(root, os.path.join(
        root, RAWS, "wellcome_global_monitor_2020", "wgmdata-covid-data-dictionary-user-guide.docx"))], func="h_wgm2020")


# ============================================================================== TISP (English master questionnaire)

TISP_FREQ_EN = [("1", "Never"), ("2", "Once or twice during the year"), ("3", "Several times during the year"),
                ("4", "Once or twice a month"), ("5", "Once or twice a week"), ("6", "Almost every day"), ("7", "Once or more per day")]
_SD5 = five("Strongly disagree", "Strongly agree")
_NA5 = five("Not at all", "Very strongly")
_REG = [("1", "Africa"), ("2", "Asia"), ("3", "Australia and Oceania"), ("4", "Europe"), ("5", "Latin America"), ("6", "North America")]
TISP_EN = {}
_S = "The following questions are about your media use. Over the past 12 months, how often have you come across information about science in the following places? "
for k, v in {"newspapersmags": "In news articles in printed newspapers or magazines", "tvradio": "In news shows or documentaries on TV or radio",
             "newswebsitesapps": "In news articles on news websites or in news apps", "videospodcasts": "In videos or podcasts on news websites or in news apps",
             "filmsseries": "In fictional films or series (e.g., on TV, in the cinema, or on other devices)", "books": "In fictional books, comics, etc.",
             "socialmedia": "On social media (e.g., YouTube vlogs, Facebook, TikTok clips, Instagram)",
             "messengers": "In instant messaging conversations with friends or family (e.g., WhatsApp, Line, Telegram)",
             "museumszoos": "In museums, zoos, or public talks",
             "rlconversations": "In conversations with friends or family (i.e., outside the Internet and messaging apps)"}.items():
    TISP_EN[f"SCIINFO_{k}"] = (_S + v, TISP_FREQ_EN, [], "ordinal", _S)
_S = "Over the past 12 months, how often have you done the following? "
for k, v in {"conversations": "Had conversations with friends, family, or co-workers about scientific issues (e.g., climate change, vaccination, nutrition, new technologies)",
             "messengers": "Chatted in messaging apps about scientific issues (e.g., WhatsApp, Telegram, Viber, Facebook Messenger)",
             "socialmedia": "Shared or commented on social media posts about scientific issues",
             "publicprotests": "Attended public rallies or protests related to scientific issues (e.g., COVID-19 protests, “Fridays for Future” demonstrations, “March for Science”)"}.items():
    TISP_EN[f"SCIENGAGE_{k}"] = (_S + v, TISP_FREQ_EN, [], "ordinal", _S)
TISP_EN["BENEFIT_ONESELF"] = ("In your opinion, how much does scientific research benefit people like yourself in your country?", _NA5, [], "ordinal", None)
TISP_EN["BENEFIT_REGION_MOST"] = ("Which region do you think benefits THE MOST from the work that scientists do?", _REG, [], "nominal", None)
TISP_EN["BENEFIT_REGION_LEAST"] = ("Which region do you think benefits THE LEAST from the work that scientists do?", _REG, [], "nominal", None)
_GOALS = {"health": "Improving public health", "energy": "Solving energy problems (renewable sources, energy security)",
          "poverty": "Reducing poverty", "defense": "Developing defense and military technology"}
for k, v in _GOALS.items():
    TISP_EN[f"GOALS_PRIO_{k}"] = ("What goals should scientists prioritize? " + v, five("Very low priority", "Very high priority"), [], "ordinal", "What goals should scientists prioritize?")
for k, v in _GOALS.items():
    TISP_EN[f"GOALS_TACKLE_{k}"] = ("How strongly do you believe that science aims to tackle these goals? " + v, _NA5, [], "ordinal", "How strongly")
_S = "The following questions are about the relationship of science and politics. To what extent do you agree or disagree with the following statements? "
for k, v in {"integrate": "Scientists should work closely with politicians to integrate scientific results into policy-making.",
             "advocate": "Scientists should actively advocate for specific policies.", "communicate": "Scientists should communicate their findings to politicians.",
             "involved": "Scientists should be more involved in the policy-making process.", "outreach": "Scientists should communicate about science with the general public.",
             "independent": "Scientists should remain independent from the policy-making process."}.items():
    TISP_EN[f"NORMPERC_{k}"] = (_S + v, _SD5, [], "ordinal", _S)
_S = "We are now interested in your opinion about personal and political decisions related to science (e.g., on energy policy, health policy). "
TISP_EN["WILLVUL_ppl"] = (_S + "How much or little should people rely on scientists’ guidance when making lifestyle choices related to science?", _NA5, [], "ordinal", _S)
TISP_EN["WILLVUL_gov"] = (_S + "How much or little should governments rely on scientists’ guidance when making decisions related to science?", _NA5, [], "ordinal", _S)
TISP_EN["WILLVUL_you"] = (_S + "How much or little control do you want scientists to have on government decisions related to science?", five("No control at all", "Very strong control"), [], "ordinal", _S)
_S = "We are interested in your opinion about scientists in your country, including scientists working at universities, in government, at companies, and for non-profit organizations. "
for k, (q, labs) in {"expert": ("How expert or inexpert are most scientists?", ["Very inexpert", "Somewhat inexpert", "Neither expert nor inexpert", "Somewhat expert", "Very expert"]),
                     "honest": ("How honest or dishonest are most scientists?", ["Very dishonest", "Somewhat dishonest", "Neither honest nor dishonest", "Somewhat honest", "Very honest"]),
                     "concerned": ("How concerned or not concerned are most scientists about people’s wellbeing?", ["Not concerned", "Somewhat not concerned", "Neither concerned nor not concerned", "Somewhat concerned", "Very concerned"]),
                     "open": ("How open are most scientists to feedback?", ["Not open", "Somewhat not open", "Neither open nor not open", "Somewhat open", "Very open"]),
                     "intellig": ("How intelligent or unintelligent are most scientists?", ["Very unintelligent", "Somewhat unintelligent", "Neither intelligent nor unintelligent", "Somewhat intelligent", "Very intelligent"]),
                     "ethical": ("How ethical or unethical are most scientists?", ["Very unethical", "Somewhat unethical", "Neither ethical nor unethical", "Somewhat ethical", "Very ethical"]),
                     "improve": ("How eager or uneager are most scientists to improve others’ lives?", ["Very uneager", "Somewhat uneager", "Neither eager nor uneager", "Somewhat eager", "Very eager"]),
                     "trans": ("How willing or unwilling are most scientists to be transparent?", ["Very unwilling", "Somewhat unwilling", "Neither willing nor unwilling", "Somewhat willing", "Very willing"]),
                     "qualified": ("How qualified or unqualified are most scientists when it comes to conducting high-quality research?", ["Very unqualified", "Somewhat unqualified", "Neither qualified nor unqualified", "Somewhat qualified", "Very qualified"]),
                     "sincere": ("How sincere or insincere are most scientists?", ["Very insincere", "Somewhat insincere", "Neither sincere nor insincere", "Somewhat sincere", "Very sincere"]),
                     "otherint": ("How considerate or inconsiderate are most scientists of others’ interests?", ["Very inconsiderate", "Somewhat inconsiderate", "Neither considerate nor inconsiderate", "Somewhat considerate", "Very considerate"]),
                     "otherviews": ("How much or little attention do scientists pay to others' views?", ["Very little attention", "Somewhat little attention", "Neither much nor little attention", "Somewhat much attention", "Very much attention"])}.items():
    TISP_EN[f"TRUST_SCI_{k}"] = (_S + q, [(str(i + 1), l) for i, l in enumerate(labs)], [], "ordinal", _S)
TISP_EN["TRUST_METHOD"] = ("To what extent do you agree or disagree with the following statement? “Scientific research methods are the best way to find out if something is true or false.”", _SD5, [], "ordinal", None)
TISP_EN["TRUST_PEW"] = ("How much confidence do you have in scientists to act in the best interests of the public?", five("No confidence at all", "A great deal of confidence"), [], "ordinal", None)
_S = "Please indicate how strongly you agree or disagree with the following statements. "
for k, v in {"othersthink": "I will share my opinions about scientific issues, regardless of what others think of them.",
             "isolate": "I will share my opinions about scientific issues, even if this may isolate me from others.",
             "against": "I will share my opinions about scientific issues, even if I believe others are against them."}.items():
    TISP_EN[f"OUTSPOKEN_{k}"] = (_S + v, _SD5, [], "ordinal", _S)
_S = "The following statements are about the relationship between science and society. How much do you agree or disagree with them? "
for k, v in {"common": "Ordinary people have in common that they trust their common sense in everyday life.", "good": "Ordinary people are of good and honest character.",
             "advantage": "Scientists are only interested in their own advantage.", "cahoots": "Scientists are in cahoots with politicians and businesses.",
             "influence": "Ordinary people should have influence on the work of scientists.", "involved": "Ordinary people should be involved in decisions about the topics scientists research.",
             "lifeexp": "Ordinary people should trust their life experience more than the recommendations of scientists.",
             "rely": "Our society should rely more on common sense than on scientific studies."}.items():
    TISP_EN[f"SCIPOP_{k}"] = (_S + v, _SD5, [], "ordinal", _S)
_S = "There are many kinds of groups in the world: men and women, ethnic and religious groups, nationalities, political factions. How much do you favor or oppose the ideas about groups in general? "
_SDO = [("1", "Extremely oppose")] + [(str(i), str(i)) for i in range(2, 10)] + [("10", "Extremely favor")]
for k, v in {"allgroupsconsider": "In setting priorities, we must consider all groups.", "notpushequality": "We should not push for group equality.",
             "equalityideal": "Group equality should be our ideal.", "superiordominate": "Superior groups should dominate inferior groups."}.items():
    TISP_EN[f"SDO_{k}"] = (_S + v, _SDO, [], "ordinal", _S)
TISP_EN["CLIM_TRUST"] = ("In the remainder of the survey, we would like to ask you some questions on the topic of climate change. To what extent do you trust scientists in your country who work on climate change?", _NA5, [], "ordinal", None)
_S = "To what extent does climate change make you feel any of the following? "
for k in ["helpless", "anxious", "optimistic", "angry", "guilty", "ashamed", "depressed", "pessimistic", "indifferent"]:
    TISP_EN[f"CLIM_EMO_{k}"] = (_S + k.capitalize(), _NA5, [], "ordinal", _S)
_S = "To what extent do you agree or disagree with the following statements? In relation to climate change I believe that my government is… "
for k, v in {"concerns": "Taking my concerns seriously", "doingenough": "Doing enough to avoid climate change", "dismisspeople": "Dismissing people’s distress",
             "science": "Acting in line with climate science", "futuregens": "Protecting future generations", "trustworthy": "Trustworthy",
             "lying": "Lying about the effectiveness of their actions taking"}.items():
    TISP_EN[f"CLIM_GOV_{k}"] = (_S + v, _SD5, [], "ordinal", _S)
_S = ("Many countries have introduced policies to reduce carbon emissions and mitigate climate change. This can include the implementation of laws aiming to "
      "reduce greenhouse gases, for example. Please indicate your level of support for the following policies. ")
for k, v in {"fueltax": "Raising carbon taxes on gas and fossil fuels or coal", "publictransport": "Expanding infrastructure for public transportation",
             "sustenergy": "Increasing the use of sustainable energy such as wind and solar energy", "protection": "Protecting forested and land areas",
             "foodtax": "Increasing taxes on carbon intense foods (e.g., beef and dairy products)"}.items():
    TISP_EN[f"CLIM_POLSUPPORT_{k}"] = (_S + v, [("1", "Not at all"), ("2", "Moderately"), ("3", "Very much")], ["4"], "ordinal", _S)
_W = {"floods": "Floods", "heatwaves": "Heatwaves", "heavystorms": "Heavy storms", "wildfires": "Wildfires", "heavyrain": "Heavy rain", "droughts": "Droughts"}
_S = ("The next questions are about climate change and weather events. When you answer them, please think about your country. To what extent do you think "
      "that climate change has increased the impact of the following weather events over the last decades? ")
for k, v in _W.items():
    TISP_EN[f"CLIM_WEATHERPAST_{k}"] = (_S + v, five("Not at all", "Very much"), [], "ordinal", _S)
_S = "(Please think about your country.) To what extent do you think that climate change will increase the impact of the following weather events in the future? "
for k, v in _W.items():
    TISP_EN[f"CLIM_WEATHERFUTU_{k}"] = (_S + v, five("Not at all", "Very much"), [], "ordinal", _S)
_S = "Political orientations are often classified on a left-right spectrum or a liberal-conservative spectrum. Please indicate your political orientation. "
TISP_EN["DEM_POL_conservative"] = (_S + "(liberal - conservative)", five("Strongly liberal", "Strongly conservative"), ["99"], "ordinal", _S)
TISP_EN["DEM_POL_right"] = (_S + "(left - right)", five("Strongly left-leaning", "Strongly right-leaning"), ["99"], "ordinal", _S)
TISP_EN["DEM_RELIGIOUS"] = ("Please indicate to what extent you consider yourself religious.", five("Not religious at all", "Very strongly religious"), [], "ordinal", None)


def h_tisp(root):
    sid = "tisp"
    base = os.path.join(root, RAWS, "osf_5c3qd_tisp")
    path = os.path.join(base, "ds_final.csv")
    man = man_sprint(root, "osf_5c3qd_tisp")
    use = ["COUNTRY_CODE", "WEIGHT_CNTRY", "DATE", "DEM_AGE", "DEM_GENDER_male", "DEM_EDU", "DEM_RESIDENCE", "UserLanguage"] + list(TISP_EN)
    df = pd.read_csv(path, sep=";", decimal=",", usecols=use, low_memory=False)
    n0 = len(df)
    cty = df.COUNTRY_CODE.map(ISO3)
    age = pd.to_numeric(df.DEM_AGE, errors="coerce")
    drop = cty.isna() | (age < 17)
    df, cty, age = df[~drop].reset_index(drop=True), cty[~drop].reset_index(drop=True), age[~drop].reset_index(drop=True)
    demo = new_demo(len(df))
    demo["country"] = cty.values
    demo["weight"] = df.WEIGHT_CNTRY.values
    demo["age_band"] = age.map(age_band)
    demo["gender"] = mf(df.DEM_GENDER_male, 1, 0)
    demo["edu3"] = pd.to_numeric(df.DEM_EDU, errors="coerce").map({1.0: "low", 2.0: "low", 3.0: "medium", 4.0: "high"})
    demo["urban"] = pd.to_numeric(df.DEM_RESIDENCE, errors="coerce").map({0.0: "rural", 1.0: "urban"})
    demo["wave"] = demo.country.map(country_waves(demo.country, ym_series(df.DATE), "2023-01"))
    ctx = L.Ctx(sid, "osf_5c3qd_tisp")
    for v, r in [("BENEFIT_OPEN, TRUST_OPEN", "open-ended free text"), ("ATTCHECK_NUMBER, ATTCHECK_RES, CONSENT", "attention checks / consent"),
                 ("DEM_GENDER*, DEM_AGE*, DEM_EDU*, DEM_RESIDENCE, DEM_INCOME*", "demographics"),
                 ("LAB, COUNTRY_*, WEIGHT_*, Start, End, DATE, Response_TYPE, Progress, Duration, Recorded, ID_*, UserLanguage", "ID / weight / admin")]:
        ctx.exclude(v, r)
    for var, (text, opts, dk, qtype, stem) in TISP_EN.items():
        add_item(ctx, var, map_codes(df[var], opts, dk), text, opts, nonsub=dk, qtype=qtype, stem=stem,
                 text_source="TISP core questionnaire (core-questionnaire_english.pdf, English master); codes as in ds_final.csv")
    langs = df.UserLanguage.astype(str).value_counts()
    meta = {"title": "TISP — Trust in Science and Science-Related Populism (Many Labs, 68 countries)", "family": "tisp",
            "source": "TISP consortium, OSF 5c3qd ds_final.csv; Mede et al. 2025, Scientific Data, doi:10.1038/s41597-024-04100-7",
            "doi_or_url": "https://osf.io/5c3qd", "licence": man.get("licence"), "commercialUse": "yes", "aiTraining": "unclear",
            "waveApprox": False, "mode": "Online self-completion (Qualtrics), national online panels by polling companies",
            "sampling": "Non-probability online quota samples (age x gender quotas), one or more labs per country",
            "weightVar": "WEIGHT_CNTRY (country-level raking weight to age/gender/education margins)",
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} (aged <17 or unmapped country code).",
                      "wave = earliest DATE month per country (>=1% of cases).",
                      "Wording = English master questionnaire; respondents answered national translations "
                      f"({len(langs)} UserLanguage codes).",
                      "SDO: master scale 1-10 (end labels); the Greek questionnaires offered 1,3,...,10 (no code 2).",
                      "CLIM_POLSUPPORT code 4 'Not applicable' -> empty; DEM_POL_* code 99 'I don't know' -> empty.",
                      "commercialUse 'yes' per CC BY 4.0 (attribution); aiTraining not stated."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path), src(root, os.path.join(base, "core-questionnaire_english.pdf"))], func="h_tisp")


# ============================================================================== COVIDiSTRESS

def h_covidistress(root):
    sid = "covidistress"
    path = os.path.join(root, RAWS, "osf_z39us_covidistress_2020", "covidistress_2020-05-30_final_cleaned.csv")
    man = man_sprint(root, "osf_z39us_covidistress_2020")
    items = list(H.COVIDISTRESS_EL)
    use = ["Country", "Dem_age", "Dem_gender", "Dem_edu", "RecordedDate", "UserLanguage"] + items
    df = pd.read_csv(path, encoding="latin-1", usecols=use, low_memory=False)
    n0 = len(df)
    cty = df.Country.map(iso2)
    age = pd.to_numeric(df.Dem_age, errors="coerce")
    drop = cty.isna() | (age < 17)
    unm = int(cty.isna().sum())
    df, cty, age = df[~drop].reset_index(drop=True), cty[~drop].reset_index(drop=True), age[~drop].reset_index(drop=True)
    demo = new_demo(len(df))
    demo["country"] = cty.values
    demo["weight"] = 1.0
    demo["age_band"] = age.map(age_band)
    demo["gender"] = df.Dem_gender.map({"Male": "male", "Female": "female"})
    demo["edu3"] = df.Dem_edu.astype(str).str.strip().map({
        "PhD/Doctorate": "high", "College degree, bachelor, master": "high",
        "Some College, short continuing education or equivalent": "medium", "Up to 12 years of school": "medium",
        "Up to 9 years of school": "low", "Up to 6 years of school": "low", "None": "low"})
    demo["wave"] = demo.country.map(country_waves(demo.country, ym_series(df.RecordedDate), "2020-04"))
    recode = {"AD_gain": {"Program A": 1, "Program B": 2}, "AD_loss": {"Program C": 1, "Program D": 2},
              "Dem_islolation": {"Life carries on as usual": 1, "Life carries on with minor changes": 2, "Isolated": 3,
                                 "Isolated in medical facility of similar location": 4}}
    ctx = L.Ctx(sid, "osf_z39us_covidistress_2020")
    for v, r in [("PSS10_avg, SLON3_avg, neu, ext, ope, agr, con, SPS_avg, Scale_UCLA_TRI_avg", "producer-derived scale scores"),
                 ("AD_check", "recognition check of the Asian-disease vignette"), ("Expl_Distress_txt, Expl_coping_txt, Final_open", "free text"),
                 ("born_92, experience_war*, war_injury, loss_during_war, time_spent_in_war*, Scale_UCLA_TRI_*, PS_PTSD_*",
                  "war-trauma supplement (country-specific module; trauma = health data)"),
                 ("Dem_riskgroup", "health-risk-group status (GDPR art. 9 health)"),
                 ("Dem_age, Dem_gender, Dem_edu, Dem_edu_mom, Dem_employment, Dem_Expat, Dem_state, Dem_maritalstatus, Dem_dependents, "
                  "Dem_isolation_adults, Dem_isolation_kids", "demographics / household counts"),
                 ("Unnamed: 0, Duration..in.seconds., RecordedDate, UserLanguage, Country", "ID / admin / timestamps")]:
        ctx.exclude(v, r)
    for var in items:
        w = H.COVIDISTRESS_EL[var]
        s = df[var].map(recode[var]) if var in recode else df[var]
        dk = [c for c, _ in w["dk"]]
        qtype = "binary" if len(w["options"]) == 2 else ("nominal" if var == "Dem_islolation" else "ordinal")
        add_item(ctx, var, map_codes(s, w["options"], dk), w["text"], w["options"], nonsub=dk, qtype=qtype, language="el",
                 stem=w["text"].split("...")[0].strip() if "..." in w["text"] else "(battery introduction included in the text)",
                 split="subsample" if var in ("AD_gain", "AD_loss") else "all",
                 text_source=f"COVIDiSTRESS v1.1 Qualtrics file (OSF z39us, file 8ex5h), Greek (EL) translation of question {w['qsf']} "
                             "(verified wording in harmonize_p4_holdouts.COVIDISTRESS_EL; English master not available offline)")
    langs = df.UserLanguage.astype(str).nunique()
    meta = {"title": "COVIDiSTRESS global survey (2020)", "family": "covidistress",
            "source": "COVIDiSTRESS consortium, OSF z39us, final cleaned file of 30 May 2020", "doi_or_url": "https://osf.io/z39us",
            "licence": man.get("licence"), "commercialUse": "yes", "aiTraining": "allowed", "waveApprox": False,
            "mode": "Online self-completion (Qualtrics), open link via researchers / social media",
            "sampling": "Convenience (snowball / social-media) samples of adults 18+, by country of residence", "weightVar": None,
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} ({unm} with country 'other'/missing; rest aged <17).",
                      "No weights: weight = 1.", "country = 'Country' (country of residence; expats are in their residence country).",
                      "wave = earliest RecordedDate month per country (>=1%).",
                      f"Item wording = Greek (EL) translation from the producers' Qualtrics file; respondents answered in {langs} "
                      "UserLanguage versions. Codes are the producers' numeric codes, common to all languages.",
                      "AD_gain / AD_loss: Asian-disease framing split ballot (each respondent saw one) -> split 'subsample'.",
                      "Expl_Distress code 99 'does not apply to my situation' -> empty. CC0 1.0 -> commercial use / AI training allowed."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path)], func="h_covidistress")


# ============================================================================== LiTS IV

LITS_EXCL = [
    (r"^q30[1-5]|^q307|^q31\d|^q32[0-5]|^q33[5-9]|^q34\d|^q350", "employment status / job characteristics / occupation history (labour-market background)"),
    (r"^q50[1-4]|^q5042", "residence / migration history (geography; demographic background)"),
    (r"^q51[0-2]$", "demographics of the household member who left"),
    (r"^q519a|^q521a|^q525a", "currency of an open remittance amount (admin detail)"),
    (r"^q701$", "employment status before the pandemic (demographic background)"),
    (r"^q713a$|^q436a$|^q438$", "household composition (demographic)"),
    (r"^q7071", "COVID-19 infection / hospitalisation in household (GDPR art. 9 health data)"),
    (r"^q80[1-4]", "health status / body measures (GDPR art. 9 health data)"),
    (r"^q8072|^q8073", "COVID-19 vaccination status and reasons (GDPR art. 9 health data)"),
    (r"^q808", "functional difficulties (disability, GDPR art. 9)"),
    (r"^q809", "anxiety / depression symptoms (GDPR art. 9 health data)"),
    (r"^q81[1-3]", "ethnic / national identity, mother tongue (GDPR art. 9 / demographic; country-specific lists)"),
    (r"^q814$", "religious denomination (GDPR art. 9 / demographic)"),
    (r"^q815|^q818$", "parents' occupation / childhood home (socio-economic background)"),
]


def parse_lits_stems(pdf):
    try:
        txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True, check=True, timeout=120).stdout
    except Exception:  # noqa: BLE001
        return {}
    lines = txt.splitlines()
    stems = {}
    for i, ln in enumerate(lines):
        mm = re.match(r"^\s*(\d)\.(\d{2})([a-z]?)\)\s+(.+)$", ln)
        if not mm:
            continue
        parts = [mm.group(4).strip()]
        for nxt in lines[i + 1:i + 8]:
            t = nxt.strip()
            if not t or re.match(r"^-?\d+\.\s", t) or re.match(r"^[a-z]\.\s", t):
                break
            letters = [c for c in t if c.isalpha()]
            if letters and sum(c.isupper() for c in letters) / len(letters) > 0.7 and not re.search(r"COUNTRY\]|^\[?\s*SURVEY", t):
                break
            parts.append(t)
        text = re.sub(r"\s+", " ", " ".join(parts))
        text = re.sub(r"\b(READ OUT|SINGLE CODE|SHOW ?CARD|MULTIPLE CODES?|INTERVIEWER)\b.*$", "", text).strip()
        text = re.sub(r"(?<=[?.:…])\s+\d{1,2}$", "", text).strip()  # page numbers caught by the layout extraction
        stems[f"{mm.group(1)}.{mm.group(2)}{mm.group(3)}"] = text
    return stems


def lits_text(var, label, stems):
    mm = re.match(r"^q(\d)(\d\d)(.*)$", var)
    if not mm:
        return label, None
    key, rest = f"{mm.group(1)}.{mm.group(2)}", mm.group(3)
    if rest[:1].isalpha() and f"{key}{rest[0]}" in stems:
        stem, sub = stems[f"{key}{rest[0]}"], rest[1:]
    elif key in stems:
        stem, sub = stems[key], rest
    else:
        return label, None
    if not sub:
        return stem, None
    return f"{stem} — {label}", stem


def h_lits4(root):
    sid = "lits4"
    base = os.path.join(root, RAWS, "ebrd_lits_iv")
    zpath = os.path.join(base, "lits_iv_dta.zip")
    man = man_sprint(root, "ebrd_lits_iv")
    stems = parse_lits_stems(os.path.join(base, "LITS-IV-Questionnaire.pdf"))
    with tempfile.TemporaryDirectory(prefix="lits4_") as tmp:
        with zipfile.ZipFile(zpath) as z:
            z.extract("lits_iv.dta", tmp)
        dta = os.path.join(tmp, "lits_iv.dta")
        _, m = pyreadstat.read_dta(dta, metadataonly=True)
        labs, vls = m.column_names_to_labels or {}, m.variable_value_labels or {}
        ctx = L.Ctx(sid, "ebrd_lits_iv")
        cand = []
        for var in m.column_names:
            lab = L.clean(labs.get(var) or "")
            if not re.match(r"^q[3-8]\d\d", var):
                continue
            if var.endswith("_skip") or not vls.get(var):
                continue
            r = next((rr for rx, rr in LITS_EXCL if re.search(rx, var)), None) or skip_reason(var, lab, sensitive=False, demographic=False)
            if r:
                ctx.exclude(var, r)
                continue
            cand.append(var)
        ctx.exclude("hhid, interviewer/visit paradata, ecs*, section 1 (household roster), section 2 (household; answered by "
                    "the knowledgeable respondent, not the sampled respondent whose demographics are recorded), i1, i2, language, "
                    "int_gender, psu, weights", "admin / roster / household-informant section / weights")
        roster = [f"q105{c}" for c in range(1, 21)] + [f"q103{c}" for c in range(1, 21)]
        use = ["country", "weight", "weight_trt", "rand_resp_code", "q109b", "urbanity", "section3_date", "start_date"] + \
              [c for c in roster if c in m.column_names] + cand
        df, _ = pyreadstat.read_dta(dta, usecols=use)
    n0 = len(df)
    clab = vls.get("country", {})
    cty = pd.to_numeric(df.country, errors="coerce").map(lambda v: iso2(clab.get(v)) if v == v else None)
    code = pd.to_numeric(df.rand_resp_code, errors="coerce")
    age = pd.Series(np.nan, index=df.index)
    sex = pd.Series(np.nan, index=df.index)
    for c in range(1, 21):
        msk = code == c
        if f"q105{c}" in df:
            age = age.where(~msk, pd.to_numeric(df[f"q105{c}"], errors="coerce"))
        if f"q103{c}" in df:
            sex = sex.where(~msk, pd.to_numeric(df[f"q103{c}"], errors="coerce"))
    age = age.where(age >= 0)
    drop = cty.isna() | (age < 17)
    keep = ~drop
    df, cty, age, sex = (x[keep].reset_index(drop=True) for x in (df, cty, age, sex))
    demo = new_demo(len(df))
    demo["country"] = cty.values
    w = pd.to_numeric(df.weight, errors="coerce")
    wt = pd.to_numeric(df.weight_trt, errors="coerce")
    use_trt = (demo.country == "TR") & wt.notna()
    demo["weight"] = w.where(~use_trt, wt).values
    demo["age_band"] = age.map(age_band)
    demo["gender"] = sex.map({1.0: "male", 2.0: "female"})
    demo["edu3"] = pd.to_numeric(df.q109b, errors="coerce").map(lambda v: None if v != v or v < 1 else "low" if v <= 3 else "medium" if v <= 5 else "high")
    demo["urban"] = pd.to_numeric(df.urbanity, errors="coerce").map({1.0: "urban", 2.0: "rural"})
    dates = pd.to_numeric(df.section3_date, errors="coerce").fillna(pd.to_numeric(df.start_date, errors="coerce"))
    demo["wave"] = demo.country.map(country_waves(demo.country, dates.map(ym_from_stata_ms), "2023-01"))
    n_stem = 0
    texts = {var: lits_text(var, L.clean(labs.get(var) or ""), stems) for var in cand}
    dup = pd.Series([t for t, _ in texts.values()]).value_counts()
    for var in cand:
        lab = L.clean(labs.get(var) or "")
        text, stem = texts[var]
        if stem is None and text != lab and dup.get(text, 0) > 1:
            text, stem = f"{text} — {lab}", text  # same question text for several variables: add the item label
        n_stem += text != lab
        L.add_labelled(ctx, var, df[var], lab, vl_mark(vls.get(var)), text=text, stem=stem,
                       qtype="ordinal" if re.match(r"^q(334|419|420|710|718)", var) else None,
                       text_source=("LiTS IV master questionnaire (LITS-IV-Questionnaire.pdf) question text" +
                                    (" + item label (Stata variable label)" if stem else "")) if text != lab else "Stata variable label")
    meta = {"title": "EBRD Life in Transition Survey IV (2022-23)", "family": "lits",
            "source": "EBRD Office of the Chief Economist, LiTS IV (lits_iv.dta in lits_iv_dta.zip) + master questionnaire",
            "doi_or_url": man.get("doi_or_url"), "licence": man.get("licence"), "commercialUse": "unclear", "aiTraining": "unclear",
            "waveApprox": False, "mode": "Face-to-face household survey (CAPI), random-walk / register sampling",
            "sampling": "National probability samples of households; one randomly selected adult (primary respondent) per household",
            "weightVar": "weight (final weight; weight_trt for Türkiye, which includes the top-up sample)",
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} (aged <17 or unmapped country).",
                      "Demographics are those of the randomly selected primary respondent (roster entry rand_resp_code: q103x sex, "
                      "q105x age); edu3 from q109b (no degree/primary/lower secondary = low; upper secondary/post-secondary non-"
                      "tertiary = medium; tertiary = high); urban from urbanity.",
                      "Items: sections 3-8 (primary respondent). Section 2 (household finances/dwelling) is answered by the "
                      "knowledgeable respondent and is excluded.",
                      f"Item wording: {n_stem} items with question text parsed from the master questionnaire (stem + item label), "
                      "others = Stata variable label. '[COUNTRY]' placeholders kept.",
                      "wave = earliest month of section-3 interview dates per country.",
                      "'Difficult to say', 'None of these', not-applicable codes -> empty.",
                      "Türkiye: weight_trt (covers the 89 top-up cases without 'weight').",
                      "XK (Kosovo) user-assigned code; PS = West Bank & Gaza."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, zpath), src(root, os.path.join(base, "LITS-IV-Questionnaire.pdf"))], func="h_lits4")


# ============================================================================== QoG EQI 2024 / 2021

def qog24_text():
    out = {}
    for v, (t, q) in H.QOG_TEXT.items():
        out[v] = (t.replace("Greece", "(COUNTRY)"), q)
    return out


QOG_STEM_RX = re.compile(r"^(In this survey, we define corruption.*?‘10’ being “strongly agree”\.\s|People engage in corruption for different reasons\..*?following\?\s|"
                         r"How much do you agree \(1 = strongly disagree, 10 = strongly agree\):\s|"
                         r"Do you agree, rather agree, rather disagree or disagree:\s|In the next few questions, we would like to get your "
                         r"opinions on some issues relating to the COVID-19 virus\.\s)", re.S)


def qog_stem(text):
    mm = QOG_STEM_RX.match(text)
    return mm.group(1).strip() if mm else None


def qog_country(nuts, label):
    p = str(nuts or "")[:2].upper()
    if re.fullmatch(r"[A-Z]{2}", p):
        return NUTS2ISO.get(p, p)
    return iso2(label)


def h_qog2024(root):
    sid = "qog2024"
    path = os.path.join(root, RAWS, "qog_eqi_2024", "qog_eqi_ind_24.dta")
    man = man_sprint(root, "qog_eqi_2024")
    _, m = pyreadstat.read_dta(path, metadataonly=True)
    vls = m.variable_value_labels or {}
    texts = qog24_text()
    items = [v for v in texts if v != "Q22"]
    df, _ = pyreadstat.read_dta(path, usecols=["country", "NUTS1_code", "D1", "D2", "D3", "D9", "Dweight", "PSweight_o", "enddate",
                                                "typeinterview"] + items)
    n0 = len(df)
    clab = vls.get("country", {})
    cty = pd.Series([qog_country(n, clab.get(c)) for n, c in zip(df.NUTS1_code, pd.to_numeric(df.country, errors="coerce"))])
    lab_cty = pd.to_numeric(df.country, errors="coerce").map(lambda c: iso2(clab.get(c)) if c == c else None)
    mism = int(((lab_cty.notna()) & (lab_cty != cty)).sum())
    age = pd.to_numeric(df.D2, errors="coerce").where(lambda a: a < 999)
    drop = cty.isna() | (age < 17)
    keep = ~drop
    df, cty, age = (x[keep].reset_index(drop=True) for x in (df, cty, age))
    demo = new_demo(len(df))
    demo["country"] = cty.values
    demo["weight"] = (pd.to_numeric(df.Dweight, errors="coerce") * pd.to_numeric(df.PSweight_o, errors="coerce")).values
    demo["age_band"] = age.map(age_band)
    demo["gender"] = mf(df.D1, 1, 2)
    demo["edu3"] = pd.to_numeric(df.D3, errors="coerce").map({1.0: "low", 2.0: "low", 3.0: "medium", 4.0: "high", 5.0: "high"})
    demo["urban"] = pd.to_numeric(df.D9, errors="coerce").map({1.0: "rural", 2.0: "urban", 3.0: "urban", 4.0: "urban"})
    demo["wave"] = demo.country.map(country_waves(demo.country, pd.to_numeric(df.enddate, errors="coerce").map(ym_from_stata_ms), "2023-09"))
    ctx = L.Ctx(sid, "qog_eqi_2024")
    ctx.exclude("Q22", "vote intention with country-specific party lists (251 party codes; >11 options, not a common item); "
                       "political item (vote) in any case")
    for v, r in [("D1, D2, D2RECODE, D3, D3RECODE, D4, D4RECODE, D5A, D5B, D8, D9", "demographics (D9 used for urban)"),
                 ("Poland_elec", "Poland-only wording flag"),
                 ("typeinterview, typetel, country, language, EQIregion, NUTS*_code/name, enddate, length", "admin / geography / timestamps"),
                 ("Pweight, Dweight, PSweight_a, PSweight_o, Party_W", "weights")]:
        ctx.exclude(v, r)
    for v in items:
        text, qtype = texts[v]
        L.add_labelled(ctx, v, df[v], L.clean(m.column_names_to_labels.get(v) or v), vl_mark(vls.get(v)), text=text, stem=qog_stem(text),
                       text_source="QoG EQI 2024 codebook (English master; verified wording in harmonize_p4_holdouts.QOG_TEXT, "
                                   "'Greece' -> '(COUNTRY)') + Stata value labels", qtype=qtype)
    mode = pd.to_numeric(df.typeinterview, errors="coerce")
    meta = {"title": "European Quality of Government Index (EQI) 2024 survey", "family": "qog",
            "source": "QoG Institute, University of Gothenburg (Charron, Lapuente, Bauhr); qog_eqi_ind_24.dta",
            "doi_or_url": man.get("doi_or_url"), "licence": man.get("licence"), "commercialUse": "unclear", "aiTraining": "unclear",
            "waveApprox": False, "mode": f"Mixed CATI ({int((mode == 1).sum())}) / CAWI online panel ({int((mode == 2).sum())})",
            "sampling": "Regionally stratified samples of residents 18+ in EU27 NUTS1/2 regions (equal-n regions)",
            "weightVar": "Dweight x PSweight_o (regional design weight x overall post-stratification)",
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} (aged <17 or no country).",
                      f"country from NUTS1_code prefix (EL -> GR); country label disagrees for {mism} rows (label code 12 is "
                      "'Pas de pays associé à ce code').",
                      "wave = earliest 'enddate' (interview date) month per country.",
                      "edu3 from D3 (1-2 = low, 3 = medium, 4-5 = high); urban from D9 (<10,000 = rural).",
                      "D6 (religious attendance) and D7 (social media hours) are behaviour items in the D-block; included.",
                      "Q20 (vote buying) mentions the vote -> excludedFromTransfer."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path), src(root, os.path.join(root, RAWS, "qog_eqi_2024", "codebook_eqi_24.pdf"))], func="h_qog2024")


def qog21_texts(pdf):
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True, check=True, timeout=120).stdout
    lines = txt.splitlines()
    out = {}
    for i, ln in enumerate(lines):
        mm = re.match(r"^\s*2\.3\.\d+\s+(q\d+(?:_\d)?)\s+-\s+(.+)$", ln)
        if not mm:
            continue
        parts = [mm.group(2).strip()]
        for nxt in lines[i + 1:i + 4]:
            t = nxt.strip()
            if not t or re.match(r"^(\(\d+\)|Very|Strongly|Agree|No confidence|Not at all|\d)", t):
                break
            parts.append(t)
        out[mm.group(1)] = re.sub(r"\s+", " ", " ".join(parts)).strip()
    return out


QOG21_CORR = ("In this survey, we define corruption to mean ‘the abuse of entrusted public power for private gain’. This ‘abuse’ could be by "
              "any public employee or politician and the ‘private gain’ might include money, gifts or other benefits. With this in mind, "
              "please respond to the following questions on corruption with a scale of 1 to 10, with ‘1’ being “strongly disagree” and "
              "‘10’ being “strongly agree”. ")
QOG21_Q17 = ("People engage in corruption for different reasons. Thinking about the reasons why people engage in corruption in your area, "
             "again, using the same scale of 1 to 10, with ‘1’ being “strongly disagree” and ‘10’ being “strongly agree”, how much do you "
             "agree with the following? ")
QOG21_COVID = "In the next few questions, we would like to get your opinions on some issues relating to the COVID-19 virus. "
QOG21_AGREE = "How much do you agree (1 = strongly disagree, 10 = strongly agree): "


def h_qog2021(root):
    """LOCKED: Greek rows (F2) only go to locked/. Options come from the Stata value labels (documentation), never from data."""
    sid = "qog2021"
    base = os.path.join(root, RAWS, "qog_eqi_2021")
    path = os.path.join(base, "qog_eqi_ind_21.dta")
    man = man_sprint(root, "qog_eqi_2021")
    enc = "WINDOWS-1252"
    try:
        _, m = pyreadstat.read_dta(path, metadataonly=True, encoding=enc)
    except Exception:  # noqa: BLE001
        enc = "latin1"
        _, m = pyreadstat.read_dta(path, metadataonly=True, encoding=enc)
    vls = m.variable_value_labels or {}
    texts = qog21_texts(os.path.join(base, "codebook_eqi_21.pdf"))
    items = [v for v in m.column_names if re.match(r"^q\d+(_\d)?$", v) and v != "q21"]
    df, _ = pyreadstat.read_dta(path, encoding=enc, usecols=["country", "NUTS1", "d1", "d3recode", "d7", "Dweight", "PSweight_o",
                                                              "typeinterview"] + items)
    n0 = len(df)
    clab = vls.get("country", {})
    cty = pd.to_numeric(df.country, errors="coerce").map(lambda c: iso2(clab.get(c)) if c == c else None)
    nprefix = df.NUTS1.astype(str).str[:2].str.upper().map(lambda p: NUTS2ISO.get(p, p))
    mism = int((cty.notna() & nprefix.str.fullmatch(r"[A-Z]{2}") & (nprefix != cty)).sum())
    drop = cty.isna()
    df, cty = df[~drop].reset_index(drop=True), cty[~drop].reset_index(drop=True)
    demo = new_demo(len(df))
    demo["country"] = cty.values
    demo["weight"] = (pd.to_numeric(df.Dweight, errors="coerce") * pd.to_numeric(df.PSweight_o, errors="coerce")).values
    demo["gender"] = mf(df.d1, 1, 2)
    demo["edu3"] = pd.to_numeric(df.d3recode, errors="coerce").map({1.0: "low", 2.0: "medium", 3.0: "high"})
    demo["urban"] = pd.to_numeric(df.d7, errors="coerce").map({1.0: "rural", 2.0: "urban", 3.0: "urban", 4.0: "urban"})
    demo["wave"] = "2020-10"
    ctx = L.Ctx(sid, "qog_eqi_2021")
    ctx.exclude("q21", "vote intention with country-specific party lists (226 party codes; >11 options, not a common item); political")
    for v, r in [("d1, d2quotas, d3recode, d4, recoded4, d5a, d5b, d6, d7", "demographics (d7 used for urban)"),
                 ("resp_id, typeinterview, typetel, country, language, EQIregion, NUTS1-3, postcode, length", "ID / admin / geography"),
                 ("Dweight, Pweight, PSweight_a, PSweight_o, Party_W, Party_W_truc", "weights")]:
        ctx.exclude(v, r)
    for v in items:
        base_t = texts.get(v) or L.clean(m.column_names_to_labels.get(v) or v)
        if re.match(r"^q1[4-6]$", v):
            text = QOG21_CORR + base_t
        elif v.startswith("q17_"):
            text = QOG21_Q17 + base_t
        elif v.startswith("q31") or v == "q32":
            text = QOG21_COVID + base_t
        elif v in ("q7", "q8", "q9", "q25", "q26", "q27", "q28", "q30"):
            text = QOG21_AGREE + base_t
        elif v in ("q10", "q11", "q12", "q13"):
            text = "Do you agree, rather agree, rather disagree or disagree: " + base_t
        else:
            text = base_t
        vl = vl_mark(vls.get(v))
        opts, nons = [], []
        for k in sorted(vl):
            lab = L.clean(vl[k])
            if L.NONSUB.search(L.norm(lab)):
                nons.append(code_str(k))
            elif lab:
                opts.append((code_str(k), lab))
        add_item(ctx, v, map_codes(df[v], opts, nons), text, opts, nonsub=nons, stem=qog_stem(text),
                 text_source="QoG EQI 2021 codebook (codebook_eqi_21.pdf, section 2.3; English master) + Stata value labels")
    mode = pd.to_numeric(df.typeinterview, errors="coerce")
    meta = {"title": "European Quality of Government Index (EQI) 2021 survey", "family": "qog",
            "source": "QoG Institute, University of Gothenburg (Charron, Lapuente, Annoni); qog_eqi_ind_21.dta",
            "doi_or_url": man.get("doi_or_url"), "licence": man.get("licence"), "commercialUse": "unclear", "aiTraining": "unclear",
            "waveApprox": True, "mode": f"Mixed CATI ({int((mode == 1).sum())}) / online ({int((mode == 2).sum())}) — all rows incl. GR",
            "sampling": "Regionally stratified samples of residents 18+ in EU27 NUTS1/2 regions",
            "weightVar": "Dweight x PSweight_o (design x overall post-stratification, as for 2024)",
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} without country. No exact age in the file (only d2quotas "
                      "18-29/30-49/50-64/65+, which does not nest the lake bands): age_band empty, no under-17 filter possible "
                      "(population 18+).",
                      f"country from the country label; NUTS1 prefix disagrees for {mism} rows.",
                      "No interview-date variable: wave = 2020-10 (codebook fieldwork 2020-10 to 2021-02) for every country (approx).",
                      "edu3 from d3recode (elementary = low, secondary = medium, third level = high); urban from d7 (<10,000 = rural).",
                      "Options from the Stata value labels (documentation), not from data; 99 DK/refuse -> empty."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path), src(root, os.path.join(base, "codebook_eqi_21.pdf"))], func="h_qog2021")


# ============================================================================== MAPLE

MAPLE_EXP = {"Q46", "Q47", "Q50", "Q53", "Q57", "Q58_1", "Q58_2", "Q61", "Q62_1", "Q62_2", "Q124", "Q130", "Q132", "Q133", "Q135",
             "Q136_1", "Q136_2", "Q136_3", "Q136_4", "Q138", "Q139", "Q141", "Q143", "Q145"}
MAPLE_DEMO = {"Q3", "Q4", "Q4A", "Q5", "Q5A", "Q6", "Q6A", "Q63", "Q64", "Q66", "Q68", "Q86_1", "Q86_2", "Q86_3", "Q86_4"}
MAPLE_PARTY_VARS = re.compile(r"^(Q27|Q28|Q33|Q34|Q35|Q36|Q37_\d+|Q38|Q38A|Q42|Q81|Q82_\d+|Q84|Q85_\d+|Q118|Q119|Q250|Q251|Q14|Q80_1|Q26_3)$")
MAPLE_W34_WEIGHT = {"PT": "weight1_pt", "GR": "weight1_gr", "IE": "weight2_ie", "DE": "weight2_ger"}


def maple_field_starts(docx_path):
    """{(ISO2 or '*', 'wN'): 'YYYY-MM'} from the codebook's 'Field(work) dates: dd.mm.(yyyy)-dd.mm.yyyy (Wave N)' lines."""
    import docx
    d = docx.Document(docx_path)
    out, cur = {}, "*"
    for p in d.paragraphs:
        t = p.text.strip()
        c = iso2(re.sub(r"\s*-.*$", "", t)) if len(t) < 40 else None
        if c:
            cur = c
            continue
        if re.match(r"(?i)^field(work)? dates", t):
            for mm in re.finditer(r"(\d{1,2})\.(\d{1,2})\.(\d{4})?\s*-\s*\d{1,2}\.\d{1,2}\.(\d{4})\s*\(Wave (\d)\)", t):
                out[(cur, f"w{mm.group(5)}")] = f"{int(mm.group(3) or mm.group(4)):04d}-{int(mm.group(2)):02d}"
    return out


def maple_wording(docx_path):
    import docx
    d = docx.Document(docx_path)
    out = {}
    for tb in d.tables:
        if len(tb.columns) != 5:
            continue
        for r in tb.rows:
            cells = [c.text.strip() for c in r.cells]
            names = re.findall(r"[A-Za-z][A-Za-z0-9_]*", cells[1].replace(" ", ""))
            word = re.sub(r"\s+", " ", cells[3]).strip()
            for nm in names:
                b = re.sub(r"_w\d$", "", nm)
                if word and b not in out:
                    out[b] = word
            qno = cells[0].strip()
            if word and re.fullmatch(r"Q\w+", qno) and qno not in out:  # 'Qno.' column (covers typos in the variable-name column)
                out[qno] = word
    return out


def h_maple(root, sid):
    if sid == "maple_w12":
        rel, slug, wv = os.path.join(RAWP, "harvard_10.7910_DVN_AGFZFO", "files"), "harvard_10.7910_DVN_AGFZFO", (("w1", {1, 3}), ("w2", {2, 3}))
        dta, cbk, ccol, wcol = "wave1_2_merged.dta", "Codebook_Wave1_2_merged.docx", "Country", "Wave"
    else:
        rel, slug, wv = os.path.join(RAWP, "harvard_10.7910_DVN_RCNK8E", "files"), "harvard_10.7910_DVN_RCNK8E", (("w3", {1, 3}), ("w4", {2, 3}))
        dta, cbk, ccol, wcol = "wave3_4merged_new.dta", "Codebook_Wave3_4_merged.docx", "country", "wave"
    path = os.path.join(root, rel, dta)
    _, m = pyreadstat.read_dta(path, metadataonly=True)
    labs, vls = m.column_names_to_labels or {}, m.variable_value_labels or {}
    wording = maple_wording(os.path.join(root, rel, cbk))
    ctx = L.Ctx(sid, slug)
    waves = [w for w, _ in wv]
    bases = {}
    for c in m.column_names:
        mm = re.match(r"^(.*)_(w\d)$", c)
        if mm and mm.group(2) in waves:
            bases.setdefault(mm.group(1), {})[mm.group(2)] = c
    cand, excl_seen = {}, set()

    def ex(b, reason):
        if b not in excl_seen:
            ctx.exclude(b, reason)
            excl_seen.add(b)
    for b, wcols in bases.items():
        lab = L.clean(labs.get(next(iter(wcols.values()))) or "")
        vl0 = vls.get(next(iter(wcols.values()))) or {}
        if re.search(r"_(gr|pt|ie|es|ger|be)$", b) or re.match(r"^class_", b):
            ex(b, "country-specific question (asked in one country only)")
            continue
        if b in MAPLE_DEMO:
            ex(b, "demographic / socio-economic background (Q68 used for urban)")
            continue
        if b in ("Q65_1", "Q65_2"):
            ex(b, "trade-union membership (GDPR art. 9 special category)")
            continue
        if b in MAPLE_EXP:
            ex(b, "survey-experiment treatment / outcome / manipulation check (treatment-dependent; assignment variables undocumented)")
            continue
        if b == "Q137":
            ex(b, "attention check")
            continue
        if sid == "maple_w34" and b == "Q24_5":
            ex(b, "different statement by country (abortion in IE/GR/PT, violence against women in ES; codebook) - not a common item")
            continue
        if re.match(r"^(time_|Q114_|Rd_|Competence|Experiment|Randomiser|CompetenceAttribution|WEIGHT|weight|recall_|id_|vote_|soph_|"
                    r"maplenews|TotalEU|MeanEU|EUsalience|Duration|duration|RecordedDate|date|ResponseId|Location|UserLanguage|device)", b):
            ex(b, "admin / paradata / weight / producer-derived recode or index")
            continue
        if not vl0:
            ex(b, "no value labels (option meanings undocumented / continuous)")
            continue
        labset = " ".join(str(v) for v in vl0.values())
        if re.search(r"\b(Party|Newspaper|Person|Region|Leader) \d+\b", labset) or re.match(r"^Q37_|^Q82_|^Q85_", b):
            ex(b, "option labels / referents are generic placeholders ('Party k', 'Leader k', 'Newspaper k') whose meaning differs "
                  "by country (party / leader / newspaper lists are country-specific)")
            continue
        cand[b] = wcols
    # read
    demo_vars = []
    for w in waves:
        demo_vars += [f"Q3_{w}", f"Q4_{w}", f"Q5_{w}", f"Q68_{w}"]
        if sid == "maple_w12":
            demo_vars += [f"WEIGHT_RIM_{w}", f"RecordedDate_{w}"]
        else:
            demo_vars += [f"Q5_es_{w}", f"date_{w}"] + [f"{p}_{w}" for p in MAPLE_W34_WEIGHT.values()]
    use = [ccol, wcol] + [c for c in demo_vars if c in m.column_names] + [c for wc in cand.values() for c in wc.values()]
    df, _ = pyreadstat.read_dta(path, usecols=list(dict.fromkeys(use)))
    n0 = len(df)
    df["_src"] = np.arange(len(df))
    blocks = []
    for w, codes in wv:
        sub = df[pd.to_numeric(df[wcol], errors="coerce").isin(list(codes))].copy()
        blk = pd.DataFrame({"_src": sub["_src"].values, "_w": w, "_cty": sub[ccol].map(iso2).values})
        blk["_age"] = pd.to_numeric(sub[f"Q4_{w}"], errors="coerce").values
        blk["_sex"] = pd.to_numeric(sub[f"Q3_{w}"], errors="coerce").map({1.0: "female", 2.0: "male"}).values
        ed = pd.to_numeric(sub[f"Q5_{w}"], errors="coerce").map(lambda v: None if v != v else "low" if v <= 3 else "medium" if v <= 5 else "high" if v <= 9 else None)
        if f"Q5_es_{w}" in sub:
            ed_es = pd.to_numeric(sub[f"Q5_es_{w}"], errors="coerce").map(lambda v: None if v != v else "low" if v <= 3 else "medium" if v == 4 else "high" if v <= 8 else None)
            ed = ed.where(ed.notna(), ed_es)
        blk["_edu"] = ed.values
        blk["_urb"] = pd.to_numeric(sub[f"Q68_{w}"], errors="coerce").map({1.0: "rural", 2.0: "urban", 3.0: "urban", 4.0: "urban"}).values
        if sid == "maple_w12":
            blk["_wt"] = pd.to_numeric(sub[f"WEIGHT_RIM_{w}"], errors="coerce").values
            blk["_ym"] = ym_series(sub[f"RecordedDate_{w}"]).values
        else:
            wt = pd.Series(np.nan, index=sub.index)
            for c2, p in MAPLE_W34_WEIGHT.items():
                if f"{p}_{w}" in sub:
                    msk = sub[ccol].map(iso2) == c2
                    wt = wt.where(~msk, pd.to_numeric(sub[f"{p}_{w}"], errors="coerce"))
            blk["_wt"] = wt.values
            blk["_ym"] = pd.to_numeric(sub[f"date_{w}"], errors="coerce").map(ym_from_stata_ms).values
        for b, wcols in cand.items():
            blk[b] = sub[wcols[w]].values if w in wcols else np.nan
        blocks.append(blk)
    st = pd.concat(blocks, ignore_index=True)
    drop = st._cty.isna() | (st._age < 17)
    st = st[~drop].reset_index(drop=True)
    demo = new_demo(len(st))
    demo["country"] = st._cty.values
    demo["weight"] = st._wt.values
    demo["age_band"] = st._age.map(age_band)
    demo["gender"] = st._sex.values
    demo["edu3"] = st._edu.values
    demo["urban"] = st._urb.values
    fs = maple_field_starts(os.path.join(root, rel, cbk))
    cw, wave_src = {}, {}
    for (c, w), g in st.groupby(["_cty", "_w"]):
        cb_start, data_m = fs.get((c, w)), L.earliest_month(g._ym.dropna())
        cw[(c, w)] = cb_start or data_m or fs.get(("*", w))
        wave_src[f"{c}/{w}"] = ("codebook country field start" if cb_start else "earliest interview month in data" if data_m
                                else "codebook all-country field start")
    demo["wave"] = [cw.get((c, w)) for c, w in zip(st._cty, st._w)]
    ctx.notes.append("wave per country x wave: " + json.dumps({k: f"{cw[tuple(k.split('/'))]} ({v})" for k, v in wave_src.items()}))
    pid = [hashlib.sha1(f"{sid}:{s}:maple-pid".encode()).hexdigest()[:12] for s in st._src]
    for b, wcols in cand.items():
        c0 = next(iter(wcols.values()))
        lab = L.clean(labs.get(c0) or "")
        vl = vls.get(c0) or {}
        for w2, c2 in wcols.items():
            if vls.get(c2) and {k: L.norm(str(v)) for k, v in vls[c2].items()} != {k: L.norm(str(v)) for k, v in vl.items()}:
                vl = {**vls[c2], **vl}  # union; documented in notes
                ctx.notes.append(f"{b}: value labels differ between waves; union of labels used")
                break
        text = L.clean(wording.get(b) or lab)
        L.add_labelled(ctx, b, st[b], lab, vl_mark(vl), text=text, stem=text if wording.get(b) else None,
                       text_source="MAPLE codebook 'Wording' column (English master)" if wording.get(b) else "Stata variable label")
        if ctx.items and ctx.items[-1]["var"] == b:
            if MAPLE_PARTY_VARS.match(b):
                tag_last(ctx, political=True, politicalReason="political protection: MAPLE party / vote / turnout / party-ID / leader item")
            only = [w for w in waves if w not in wcols]
            if only:
                tag_last(ctx, note=f"asked only in wave(s) {sorted(wcols)}; rows of other waves are empty")
    meta_w12 = sid == "maple_w12"
    meta = {"title": "MAPLE " + ("I panel, waves 1-2 (2019)" if meta_w12 else "II panel, waves 3-4 (2019-2021)"),
            "family": "maple", "source": f"MAPLE project (Harvard Dataverse doi:10.7910/DVN/{'AGFZFO' if meta_w12 else 'RCNK8E'}), {dta}",
            "doi_or_url": f"https://doi.org/10.7910/DVN/{'AGFZFO' if meta_w12 else 'RCNK8E'}",
            "licence": "Harvard Dataverse terms (deposit labelled CC0 1.0 per harmonize.py metadata; verify on Dataverse)",
            "commercialUse": "unclear", "aiTraining": "unclear", "waveApprox": False,
            "mode": ("Online panels (Qualtrics/Netquest)" if meta_w12 else
                     "Online panels (Qualtrics/Netquest) except Greece: representative phone survey (Metron Analysis)"),
            "sampling": "Quota samples (gender x age x education, NUTS1) of adults 18+; panel: one row per respondent-wave",
            "weightVar": ("WEIGHT_RIM_w1 / WEIGHT_RIM_w2 (RIM weights)" if meta_w12 else
                          "demographic weights: weight1_pt, weight1_gr (not weight2_gr: includes past vote), weight2_ie (rim), "
                          "weight2_ger (3 education groups); Spain: none (weight 1)"),
            "notes": [f"Rows (respondents) in file: {n0}; respondent-wave rows: {len(st) + int(drop.sum())}; dropped {int(drop.sum())} "
                      "(aged <17 / unmapped country).",
                      "One row per respondent-wave; 'pid' = salted hash of the source row (links the two waves of a respondent; "
                      "source ResponseId/ID not kept - ID contains IP-like strings).",
                      "Items = base variable names (suffix _wN dropped); items asked in one wave only are empty in the other wave.",
                      "edu3 from Q5 (ISCED-like: early childhood/primary/lower secondary = low, upper/post-secondary = medium, "
                      "tertiary = high; Spain Q5_es). urban from Q68 (rural area or village = rural; towns/suburbs/cities = urban).",
                      "All party / vote / turnout / party-ID / leader items carry excludedFromTransfer; vote-choice items with "
                      "generic 'Party k' codes and >11 options are excluded."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, path), src(root, os.path.join(root, rel, cbk))], pid=pid, func=f"h_maple({sid})")


# ============================================================================== TealHelix (adapted from harmonize_p4_lake.h_tealhelix)

def h_tealhelix(root):
    sid = "tealhelix"
    slug = "zenodo_18631009_tealhelix_food_labels_11eu"
    man = man_phase4(root, slug)
    base = os.path.join(root, RAW4, slug, "extracted", "Replication_Package_Tealhelix")
    num = pd.read_csv(os.path.join(base, "TealHelix.csv"), header=None, dtype=str, keep_default_na=False)
    txt = pd.read_csv(os.path.join(base, "TealHelix2.csv"), header=None, dtype=str, keep_default_na=False)
    assert num.shape == txt.shape
    qrow = num.iloc[0].replace("", np.nan).ffill().fillna("")
    srow = num.iloc[1]
    dn, dt_ = num.iloc[2:].reset_index(drop=True), txt.iloc[2:].reset_index(drop=True)
    n0 = len(dn)
    live = next(j for j in range(num.shape[1]) if num.iat[0, j].startswith("Which country do you currently live in?"))
    cty = dt_[live].str.strip().map(iso2)
    col_age = next(j for j in range(num.shape[1]) if num.iat[0, j].startswith("Age:"))
    age = pd.to_numeric(dt_[col_age], errors="coerce")
    drop = cty.isna() | (age < 17)
    dn, dt_, cty, age = (x[~drop].reset_index(drop=True) for x in (dn, dt_, cty, age))
    demo = new_demo(len(dn))
    demo["country"] = cty.values
    demo["weight"] = 1.0
    demo["age_band"] = age.map(age_band)
    col_g = next(j for j in range(num.shape[1]) if num.iat[0, j].startswith("Gender:"))
    demo["gender"] = dt_[col_g].str.strip().str.lower().map({"male": "male", "female": "female", "man": "male", "woman": "female"})
    col_e = next(j for j in range(num.shape[1]) if num.iat[0, j].startswith("Education:"))

    def edu(t):
        t = L.norm(t)
        if not t or "prefer" in t:
            return None
        if re.search(r"bachelor|master|doctor|phd|tertiary|university|degree", t):
            return "high"
        if re.search(r"secondary|high school|vocational|post-secondary|upper", t):
            return "medium"
        if re.search(r"primary|no formal|less than|elementary|lower", t):
            return "low"
        return None
    demo["edu3"] = dt_[col_e].map(edu)
    demo["wave"] = "2025-07"
    ctx = L.Ctx(sid, slug)
    skip_q = re.compile(r"^(Respondent ID|Collector ID|Start Date|End Date|IP Address|Email Address|First Name|Last Name|Custom Data|"
                        r"I have read the TealHelix|What is your Prolific|Age:|Which country|Please name the city|Gender:|Education:|"
                        r"How many children|How many people live|What is your current occupation|What is your average monthly household income|"
                        r"Most modern theories of decision-making|Please choose the picture below|According to the Food and Agricultural|"
                        r"Please indicate the extra percentage)")
    reasons = {"Most modern theories": "attention check (instructional manipulation check)",
               "Please choose the picture": "pictorial response options (pictures undocumented in deposit)",
               "According to the Food": "constant-sum allocation (0-100 points; continuous)",
               "Please indicate the extra percentage": "open numeric answer",
               "Which country": "screening / demographic", "Please name the city": "open text (place of residence)"}
    rng = {}
    try:
        dic = pd.read_excel(os.path.join(base, "Dictionary.xlsx"), dtype=object)
        dic.columns = ["q", "a", "mn", "mx"]
        start = next(j for j in range(num.shape[1]) if num.iat[0, j].startswith("What is your type of diet"))
        for k, r in dic.iterrows():
            c = start + k
            if c >= num.shape[1]:
                break
            if L.norm(L.fix_text(str(r.q)))[:25].replace("â€™", "'") and L.norm(str(qrow[c]))[:12] != L.norm(L.fix_text(str(r.q)))[:12]:
                continue
            m1, m2 = re.match(r"^\s*(\d+)", str(r.mn)), re.match(r"^\s*(\d+)", str(r.mx))
            if m1 and m2:
                rng[c] = (int(m1.group(1)), int(m2.group(1)))
    except Exception as e:  # noqa: BLE001
        ctx.notes.append(f"Dictionary.xlsx not usable for option ranges ({e})")
    qlabels = {}
    for c in range(num.shape[1]):
        pr = pd.DataFrame({"n": num.iloc[2:, c].str.strip(), "t": txt.iloc[2:, c].str.strip()})
        pr = pr[(pr.n.str.fullmatch(r"\d+")) & (pr.t != "")].drop_duplicates()
        for n_, t_ in zip(pr.n, pr.t):
            qlabels.setdefault(L.clean(qrow[c]), {}).setdefault(int(n_), t_)
    incomplete, done_groups = [], set()
    j, ncol = 0, num.shape[1]
    while j < ncol:
        q, sub, var = L.clean(qrow[j]), L.clean(srow[j]), f"c{j:03d}"
        if skip_q.match(q):
            key = next((v for k, v in reasons.items() if q.startswith(k)), "admin / PII header / demographic")
            if (q, key) not in done_groups:
                ctx.exclude(f"{var} ({q[:60]})", key)
                done_groups.add((q, key))
            j += 1
            continue
        if sub.startswith("http"):
            ctx.exclude(f"{var} ({q[:50]} / image)", "sub-item is an image (label logo) without documented name")
            j += 1
            continue
        if sub.startswith("Open-Ended") or re.match(r"other \(please specify\)", sub, re.I):
            ctx.exclude(f"{var} ({q[:50]})", "open text")
            j += 1
            continue
        group, k = [j], j + 1
        while k < ncol and num.iat[0, k] == "" and L.clean(qrow[k]) == q:
            group.append(k)
            k += 1
        tvals = set(dt_[j].unique()) - {""}
        is_checkbox = (sub != "Response" and len(tvals) <= 1 and (not tvals or L.clean(next(iter(tvals))) == sub)) and len(group) > 1
        if is_checkbox:
            any_ticked = dt_[group].apply(lambda col: col.str.strip().ne("")).any(axis=1)
            for c in group:
                s = L.clean(srow[c])
                if re.match(r"other \(please specify\)", s, re.I):
                    ctx.exclude(f"c{c:03d}", "open 'other' option")
                    continue
                codes = dt_[c].str.strip().ne("").map({True: "1", False: "0"}).where(any_ticked, None).astype(object)
                add_item(ctx, f"c{c:03d}", codes, f"{q} — mentions: «{s}»", [("0", "Not selected"), ("1", s)], label=f"{q}: {s}",
                         text_source="SurveyMonkey header (question + option)", multi=True, qtype="binary")
            j = k
            continue
        pairs = pd.DataFrame({"n": dn[j].str.strip(), "t": dt_[j].str.strip()})
        pairs = pairs[(pairs.n != "") & (pairs.t != "")]
        mp = pairs.drop_duplicates().groupby("n").t.agg(lambda x: sorted(set(x)))
        if pairs.empty or (mp.map(len) > 1).any() or not all(re.fullmatch(r"\d+", c) for c in mp.index):
            ctx.exclude(var, "numeric and text exports disagree / empty / non-numeric codes")
            j += 1
            continue
        opts = {int(c): ts[0] for c, ts in mp.items()}
        nonsub = {c for c, l in opts.items() if L.NONSUB.search(L.norm(l))}
        subst = {c: l for c, l in opts.items() if c not in nonsub}
        if j in rng:
            lo, hi = rng[j]
            if hi - lo <= 10 and all(lo <= c <= hi for c in subst):
                unknown = []
                for c in [c for c in range(lo, hi + 1) if c not in subst and c not in nonsub]:
                    lab_c = qlabels.get(q, {}).get(c)
                    if lab_c and L.NONSUB.search(L.norm(lab_c)):
                        nonsub.add(c)
                    else:
                        subst[c] = lab_c or str(c)
                        if not lab_c:
                            unknown.append(c)
                if unknown:
                    incomplete.append(f"{var}: codes {unknown} in the Dictionary.xlsx range {lo}-{hi} never chosen: numeric label used")
        codes = pd.to_numeric(dn[j], errors="coerce").map(lambda v: None if v != v or int(v) in nonsub else str(int(v))).astype(object)
        text = q if sub in ("Response", "") else f"{q} — {sub}"
        add_item(ctx, var, codes, text, [(str(c), subst[c]) for c in sorted(subst)], nonsub=[str(c) for c in sorted(nonsub)],
                 label=text, stem=q if sub not in ("Response", "") else None,
                 text_source="SurveyMonkey header (question + sub-item); option labels from the parallel text export")
        j += 1
    ctx.notes.extend(incomplete)
    meta = {"title": "TealHelix consumer survey on food sustainability labels (11+ EU countries)", "family": "tealhelix",
            "source": man.get("producer"), "doi_or_url": man.get("doi"), "licence": man.get("licence"), "commercialUse": "yes",
            "aiTraining": "unclear", "waveApprox": True, "mode": "Online self-completion (SurveyMonkey), Prolific panel",
            "sampling": "Non-probability Prolific quota samples (~300 per country)", "weightVar": None,
            "notes": [f"Rows in file: {n0}; dropped {int(drop.sum())} (country of residence 'Other'/blank or aged <17).",
                      "country = 'Which country do you currently live in?'. No weight (=1), no urban variable.",
                      "wave = 2025-07 APPROXIMATION (Start/End Date columns empty; project 2024-2027; export dated 2026-02-09).",
                      "edu3 parsed from the education answer text. Option labels from the parallel text export (TealHelix2.csv)."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, os.path.join(base, "TealHelix.csv")), src(root, os.path.join(base, "TealHelix2.csv")),
                                                 src(root, os.path.join(base, "Dictionary.xlsx"))], func="h_tealhelix")


# ============================================================================== Populism IT / GR / ES 2022 (adapted from h_populism)

POP_STEMS_EN = {  # block instructions from Survey_draft_in_english.pdf (English master draft)
    "pub_pol": "Policy views. Please indicate your level of agreement with the following statements:",
    "values": "Values. There are a number of qualities that people feel children should have. Below we provide pairs of desirable "
              "qualities. For each pair of qualities, please tell us which one you think is more important",
    "social_views1": "Social views 1. Please indicate your level of agreement with the following statements:",
    "pol_ideas1": "Political views 1. Please indicate your level of agreement with the following statements:",
    "pol_ideas2": "Political views 2. Please indicate your level of agreement with the following statements:",
    "pol_system1": "Political system 1. To what extent you are willing to accept... (rate from 1 \"Not at all\" to 5 \"Completely willing\" to accept)",
    "trust": "Trust. On a score of 0-10 how much you personally trust each of the following. 0 means you do not trust an institution "
             "at all, and 10 means you have complete trust",
}


def h_populism(root):
    sid = "populism2022"
    slug = "zenodo_20772797_populism_it_gr_es_2022"
    man = man_phase4(root, slug)
    d = os.path.join(root, RAW4, slug)
    g = pd.read_csv(os.path.join(d, "df_south.csv"))
    cb = pd.read_excel(os.path.join(d, "codebook_ESP_ITA_GRC.xlsx"), sheet_name="Hoja 1", dtype=object)
    n0 = len(g)
    cty = g.cntry.map(ISO3)
    g, cty = g[cty.notna()].reset_index(drop=True), cty[cty.notna()].reset_index(drop=True)
    demo = new_demo(len(g))
    demo["country"] = cty.values
    demo["weight"] = 1.0
    demo["age_band"] = pd.to_numeric(g.age_groups, errors="coerce").map({1.0: "17-24", 2.0: "25-34", 3.0: "35-44", 4.0: "45-54"})
    demo["gender"] = pd.to_numeric(g.gender, errors="coerce").map({1.0: "male", 2.0: "female"})
    demo["edu3"] = pd.to_numeric(g.education, errors="coerce").map({1.0: "low", 2.0: "low", 3.0: "medium", 4.0: "medium", 5.0: "high"})
    demo["wave"] = "2022-03"
    ctx = L.Ctx(sid, slug)
    cbd = {r.variable: r for _, r in cb.iterrows() if isinstance(r.variable, str)}
    skip_demo = {"Unnamed: 0", "age_yrs", "gender", "education", "employ_status", "domicile", "religion", "age_groups", "cntry"}
    for var in g.columns:
        if var in skip_demo:
            ctx.exclude(var, "demographic / index column" + ("; domicile: 4 data codes vs 5 documented levels" if var == "domicile" else ""))
            continue
        if re.search(r"_(esp|ita|grc)$|esp$", var):
            ctx.exclude(var, "country-specific question (party affinity / vote / neighbours / distance for one country's party list)")
            continue
        r = cbd.get(var)
        if r is None:
            ctx.exclude(var, "not in codebook")
            continue
        wording = L.clean(r.wording)
        if var == "lrscale":
            opts = [(str(k), "0 left" if k == 0 else "10 right" if k == 10 else str(k)) for k in range(11)]
            note_src = "codebook wording (end labels from the wording; codebook 'levels' is a copy-paste error)"
        elif var == "pol_system2.7":
            opts = [(str(k), str(k)) for k in range(11)]
            note_src = "codebook wording; scale end labels undocumented (codebook levels refer to satisfaction), options numeric"
        else:
            opts = L._pop_options(r.levels, r["values"], wording)
            note_src = "codebook wording + levels"
        if not opts:
            ctx.exclude(var, f"codebook levels/values not parseable ({r.levels} / {r['values']})")
            continue
        num = pd.to_numeric(g[var], errors="coerce")
        valid = {c for c, _ in opts}
        codes = num.map(code_str).astype(object)
        unl = codes.notna() & ~codes.isin(valid) & (num != 99)
        if unl.sum():
            ctx.exclude(var, f"{int(unl.sum())} answers with codes outside the documented range")
            continue
        codes = codes.where(codes.isin(valid), None)
        block = var.split(".")[0] if var.startswith(("social_views", "pol_ideas", "pol_system")) else re.sub(r"\d+$", "", var.split(".")[0])
        stem = POP_STEMS_EN.get(block)
        text = f"{stem} — {wording}" if stem and not wording.lower().startswith("how much do you personally trust") else wording
        add_item(ctx, var, codes, text, opts, nonsub=["99"] if (num == 99).any() else [], label=L.clean(r.label),
                 stem=stem if text != wording else None,
                 text_source=("English master draft block instruction (Survey_draft_in_english.pdf) + " if text != wording else "") + note_src)
    meta = {"title": "Populist and other socio-political attitudes in Italy, Greece and Spain (2022)", "family": "populism",
            "source": man.get("producer"), "doi_or_url": man.get("doi"), "licence": man.get("licence"), "commercialUse": "yes",
            "aiTraining": "unclear", "waveApprox": False, "mode": "Online self-completion (CAWI, Qualtrics panel)",
            "sampling": "Non-probability Qualtrics quota samples of adults 18+ (1,082 per country)", "weightVar": None,
            "notes": [f"Rows in df_south.csv: {n0}; none dropped (18+ sample).",
                      "age_band from age_groups (18-24 -> 17-24 ... 45-54; '55+' cannot be split -> empty); 'age_yrs' is not age in years.",
                      "No weights (=1). edu3: no formal/primary = low, secondary/professional = medium, university = high.",
                      "wave = 2022-03 (Zenodo description: fieldwork 2022-03-23 to 2022-04-05).",
                      "Item wording = English codebook wording, with English block instructions from the English draft questionnaire.",
                      "Country-specific party / vote / neighbour items (_esp/_ita/_grc) excluded: not common items."]}
    return finalize(root, sid, ctx, demo, meta, [src(root, os.path.join(d, "df_south.csv")), src(root, os.path.join(d, "codebook_ESP_ITA_GRC.xlsx")),
                                                 src(root, os.path.join(d, "Survey_draft_in_english.pdf"))], func="h_populism")


# ============================================================================== overlap check (ids / years only)

def overlap_check(root):
    out = {"generatedAt": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "checks": []}
    ids = {}
    for y in (2019, 2021, 2023, 2025):
        p = os.path.join(root, RAW4, "lrf_world_risk_poll_2019_2025", "extracted", f"WRP_{y}.sav")
        if os.path.exists(p):
            d, _ = pyreadstat.read_sav(p, usecols=["WPID_RANDOM"])
            ids[f"wrp{y}"] = set(pd.to_numeric(d.WPID_RANDOM, errors="coerce").dropna().astype("int64"))
    p = os.path.join(root, RAWS, "wellcome_global_monitor_2020", "wgm-full-wave2-public-file.sav")
    if os.path.exists(p):
        d, _ = pyreadstat.read_sav(p, usecols=["WPID_RANDOM"])
        ids["wgm2020"] = set(pd.to_numeric(d.WPID_RANDOM, errors="coerce").dropna().astype("int64"))
    keys = sorted(ids)
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            out["checks"].append({"files": [a, b], "idVariable": "WPID_RANDOM (Gallup random case id)",
                                  "nA": len(ids[a]), "nB": len(ids[b]), "sharedIds": len(ids[a] & ids[b])})
    try:
        d12, _ = pyreadstat.read_dta(os.path.join(root, RAWP, "harvard_10.7910_DVN_AGFZFO", "files", "wave1_2_merged.dta"),
                                     usecols=["ResponseId_w1", "ResponseId_w2"])
        d34, _ = pyreadstat.read_dta(os.path.join(root, RAWP, "harvard_10.7910_DVN_RCNK8E", "files", "wave3_4merged_new.dta"), usecols=["ID"])
        s12 = set(d12.ResponseId_w1.dropna().astype(str)) | set(d12.ResponseId_w2.dropna().astype(str))
        s34 = set(d34.ID.dropna().astype(str))
        out["checks"].append({"files": ["maple_w12", "maple_w34"], "idVariable": "ResponseId_w1/_w2 vs ID",
                              "nA": len(s12 - {""}), "nB": len(s34 - {""}), "sharedIds": len((s12 & s34) - {""})})
    except Exception as e:  # noqa: BLE001
        out["checks"].append({"files": ["maple_w12", "maple_w34"], "error": str(e)})
    out["reasoning"] = [
        "Gallup World Poll add-ons: wgm2018 = GWP 2018 (no case id in the public xlsx), wrp2019 = GWP 2019, wgm2020 = GWP 2020, "
        "wrp2021 / wrp2023 / wrp2025 = GWP 2021 / 2023 / 2025. Gallup draws fresh national samples every wave, so files from "
        "different years share no respondents; the WPID_RANDOM comparison above confirms this where ids exist.",
        "wgm2018 vs wrp2019: different GWP years (2018 vs 2019); different respondents (no id in wgm2018 to compare).",
        "qog2021 vs qog2024: independent cross-sections (2020-21 vs 2023-24); no shared id variable (2024 file has none).",
        "maple_w12 (MAPLE I, 2019) vs maple_w34 (MAPLE II, 2019-2021): separate panels; ids compared above. Within each file, "
        "the two waves of a panel respondent are linked by 'pid' (same person, by design).",
        "tisp (2023), covidistress (2020), lits4 (2022-23), tealhelix (2025?), populism2022: independent studies; no ids shared "
        "with any other file. Greek rows of populism2022 and maple were already used as Phase-4 dev evidence (not new holdouts).",
        "No file in this scope is a replication copy containing the locked Greek rows of wgm2018 (F1) or qog2021 (F2)."]
    with open(out_dir(root, "provenance", "_overlap.world.json"), "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    for c in out["checks"]:
        print("overlap:", c.get("files"), "shared ids:", c.get("sharedIds"), c.get("error", ""))
    return out


# ============================================================================== main

ADAPTERS = {
    "wrp2019": lambda r: h_wrp(r, 2019), "wrp2021": lambda r: h_wrp(r, 2021), "wrp2023": lambda r: h_wrp(r, 2023),
    "wrp2025": lambda r: h_wrp(r, 2025), "wgm2018": h_wgm2018, "wgm2020": h_wgm2020, "tisp": h_tisp,
    "covidistress": h_covidistress, "lits4": h_lits4, "qog2021": h_qog2021, "qog2024": h_qog2024,
    "maple_w12": lambda r: h_maple(r, "maple_w12"), "maple_w34": lambda r: h_maple(r, "maple_w34"),
    "tealhelix": h_tealhelix, "populism2022": h_populism,
}


def validate(root, sids):
    problems = []
    for sid in sids:
        rp = os.path.join(root, OUT_REL, "resp", f"{sid}.csv.gz")
        ip = os.path.join(root, OUT_REL, "items", f"{sid}.json")
        if not (os.path.exists(rp) and os.path.exists(ip)):
            problems.append(f"{sid}: missing resp/items")
            continue
        hdr = pd.read_csv(rp, nrows=0).columns.tolist()
        items = json.load(open(ip, encoding="utf-8"))
        if not items:
            problems.append(f"{sid}: no items")
        if hdr[:len(DEMO_OUT)] != DEMO_OUT:
            problems.append(f"{sid}: resp columns {hdr[:8]}")
        if [c for c in hdr if "::" in c] != [it["id"] for it in items]:
            problems.append(f"{sid}: item columns != items json")
        cc = pd.read_csv(rp, usecols=["country", "wave", "weight"], dtype={"country": str, "wave": str}, keep_default_na=False)
        cc["weight"] = pd.to_numeric(cc.weight)
        if len(cc) == 0:
            problems.append(f"{sid}: no rows")
        if not cc.country.str.fullmatch(r"[A-Z]{2}").all():
            problems.append(f"{sid}: non-ISO2 country codes")
        if not cc.wave.fillna("").str.fullmatch(r"\d{4}-\d{2}").all():
            problems.append(f"{sid}: wave not YYYY-MM for some rows")
        wm = cc.groupby("country").weight.mean()
        if ((wm - 1).abs() > 1e-3).any():
            problems.append(f"{sid}: weights not mean 1 within country")
        if sid in LOCKED:
            if (cc.country == "GR").any():
                problems.append(f"{sid}: GR rows in resp (locked!)")
            lp = os.path.join(root, OUT_REL, "locked", f"{sid}.GR.csv.gz")
            if not os.path.exists(lp) or pd.read_csv(lp, nrows=0).columns.tolist() != hdr:
                problems.append(f"{sid}: locked file missing or columns differ")
            elif len(pd.read_csv(lp, usecols=["rid"])) == 0:
                problems.append(f"{sid}: locked file has no rows")
        for it in items:
            if not 2 <= len(it["options"]) <= 11 or it["nValid"] < MIN_VALID:
                problems.append(f"{it['id']}: option/valid rule violated")
    return problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".")
    ap.add_argument("--only", action="append", choices=sorted(ADAPTERS))
    ap.add_argument("--no-overlap", action="store_true")
    ap.add_argument("--validate-only", action="store_true")
    a = ap.parse_args()
    root = os.path.abspath(a.root)
    sids = a.only or list(ADAPTERS)
    sp = os.path.join(root, OUT_REL, SURVEYS_FILE)
    entries = {}
    if os.path.exists(sp):
        entries = {e["surveyId"]: e for e in json.load(open(sp, encoding="utf-8"))}
    for sid in ([] if a.validate_only else sids):
        e = ADAPTERS[sid](root)
        entries[sid] = e
        order = list(ADAPTERS)
        with open(out_dir(root, SURVEYS_FILE), "w", encoding="utf-8") as fh:
            json.dump(sorted(entries.values(), key=lambda x: order.index(x["surveyId"]) if x["surveyId"] in order else 99), fh,
                      ensure_ascii=False, indent=1)
    if not (a.no_overlap or a.validate_only):
        overlap_check(root)
    problems = validate(root, sids)
    print("validation: OK" if not problems else "validation problems:\n  " + "\n  ".join(problems[:60]))
    if problems:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
