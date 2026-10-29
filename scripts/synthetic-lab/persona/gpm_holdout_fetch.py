"""
PHASE6 — νέο ΤΥΦΛΟ τελικό holdout (H6): Eurobarometer 2025–2026 που ΔΕΝ υπάρχουν στη λίμνη (lake.sqlite), στα PDF
του Lab ή στη διεθνή βάση της Phase 5 (έλεγχος: docs/synthetic-lab/PHASE6_PREREGISTRATION.md §2).

Ροή (κανένα ελληνικό ποσοστό δεν τυπώνεται ή διαβάζεται από άνθρωπο):
  1. API data.europa.eu → αρχείο Volume A (πίνακες ανά χώρα, σταθμισμένοι), λήψη στο .synthetic-lab-data/raw/phase6/holdout/<id>/
  2. ανάλυση: ερώτηση, επιλογές (μόνο οι βασικές γραμμές, όχι τα «Total …»), ποσοστό ανά χώρα, μη σταθμισμένη βάση
  3. η στήλη EL (Ελλάδα) γράφεται ΜΟΝΟ σε locked/<id>.EL.json (sha256 στο μανιφέστο)· οι άλλες χώρες στο visible/<id>.json
  4. Volume C (ανά χώρα × δημογραφικά): από το zip αποθηκεύεται μόνο το ελληνικό αρχείο, χωρίς ανάγνωση, στο locked/.

  python3 -I scripts/synthetic-lab/persona/gpm_holdout_fetch.py --root .
"""
import argparse
import hashlib
import io
import json
import os
import re
import subprocess
import time
import zipfile

H6 = {
    "fl015ep": "s3732_fl015ep_eng", "fl574": "s3652_fl574_eng", "eb050ep": "s3752_105_3_eb050ep_eng", "eb049ep": "s3632_104_2_eb049ep_eng",
    "sp564": "s3413_103_2_sp564_eng", "fl563": "s3355_fl563_eng", "eb048ep": "s3572_103_4_eb048ep_eng", "eb047ep": "s3492_103_1_eb047ep_eng",
    "fl570": "s3385_fl570_eng",
}
GREEK = "EL"


def sha(b):
    return hashlib.sha256(b).hexdigest()


def curl(url, out=None):
    for d in (2, 4, 8, 16):
        r = subprocess.run(["curl", "-sS", "-m", "300", "-L", "--fail"] + (["-o", out] if out else []) + [url], capture_output=True)
        if r.returncode == 0:
            return r.stdout
        time.sleep(d)
    raise RuntimeError(f"λήψη απέτυχε: {url}: {r.stderr[:200]}")


def distributions(ds):
    d = json.loads(curl(f"https://data.europa.eu/api/hub/search/datasets/{ds}"))["result"]
    out = []
    for x in d.get("distributions", []):
        t = x.get("title")
        t = (t.get("en") if isinstance(t, dict) else t) or ""
        url = (x.get("access_url") or x.get("download_url") or [None])[0]
        out.append((t, url))
    return d, out


def parse_volume_a(path):
    """Δύο μορφές: Flash (ετικέτα → γραμμή πλήθους, κενή ετικέτα → γραμμή αναλογίας) και Standard/EP (γαλλική ετικέτα →
    πλήθος, αγγλική ετικέτα → αναλογία). Κρατούνται μόνο οι βασικές επιλογές (σταματά στο πρώτο «Total …»)."""
    import openpyxl

    QRE = re.compile(r"^(Q[A-Z]?\d|D\d|[A-Z]{1,3}\d+)[A-Za-z0-9_.]*[.)\s]")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    items = []
    num = lambda v: 0.0 if v == "-" else (float(v) if isinstance(v, (int, float)) else None)
    for sn in wb.sheetnames:
        if sn.lower().startswith(("table of", "content", "note")) or sn in ("B",):
            continue
        rows = list(wb[sn].iter_rows(values_only=True))
        h = next((i for i, r in enumerate(rows) if any(c == GREEK for c in r if isinstance(c, str))), None)
        if h is None:
            continue
        hdr = {}
        for j, c in enumerate(rows[h]):
            if isinstance(c, str):
                c2 = c.split("\n")[-1].strip()
                if re.match(r"^[A-Z]{2,4}\d{0,2}$", c2):
                    hdr[c2] = j
        lc = max(min(hdr.values()) - 1, 0) if hdr else 1  # στήλη ετικετών: αμέσως πριν από την πρώτη χώρα
        for r in rows[h + 1: h + 10]:  # ή η στήλη όπου εμφανίζεται η γραμμή «Base»
            jb = next((j for j, c in enumerate(r[:4]) if isinstance(c, str) and c.startswith("Base")), None)
            if jb is not None:
                lc = jb
                break
        q, base = None, None
        for r in rows[:h]:
            for c in reversed(r):  # αγγλικό κείμενο (δεξιά) κατά προτίμηση
                if isinstance(c, str) and QRE.match(c.strip()) and q is None:
                    q = c.strip()
                if isinstance(c, str) and c.startswith("Base:") and base is None and re.search(r"[a-z]", c) and "Ensemble" not in c:
                    base = c.strip()
        opts, nb = [], {}
        i = h + 1
        while i < len(rows) - 1:
            lab = rows[i][lc] if len(rows[i]) > lc else None
            if isinstance(lab, str) and lab.strip() in ("Total", "Base: Total (unweighted)", "Base: Total"):
                nb = {c: rows[i][j] for c, j in hdr.items()}
                i += 1
                continue
            if isinstance(lab, str) and lab.startswith("Base:"):
                i += 1
                continue
            if isinstance(lab, str) and lab.strip().lower().startswith("total"):
                break
            if isinstance(lab, str) and lab.strip():
                nxt = rows[i + 1]
                nl = nxt[lc] if len(nxt) > lc else None
                vals = {c: num(nxt[j]) if j < len(nxt) else None for c, j in hdr.items()}
                if any(v is not None for v in vals.values()):
                    opts.append(((nl if isinstance(nl, str) and nl.strip() else lab).strip(), vals))
                    i += 2
                    continue
            i += 1
        if q and len(opts) >= 2:
            items.append({"sheet": sn, "question": q, "base": base, "options": [o for o, _ in opts], "values": {c: [v.get(c) for _, v in opts] for c in hdr}, "nUnweighted": nb})
    return items


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".")
    a = ap.parse_args()
    root = os.path.abspath(a.root)
    raw = os.path.join(root, ".synthetic-lab-data", "raw", "phase6", "holdout")
    out = os.path.join(root, ".synthetic-lab-data", "gpm", "holdout-h6")
    os.makedirs(os.path.join(out, "visible"), exist_ok=True)
    os.makedirs(os.path.join(out, "locked"), exist_ok=True)
    man = {"version": "h6-1", "createdAt": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "surveys": {}}
    for sid, ds in H6.items():
        d, dist = distributions(ds)
        os.makedirs(os.path.join(raw, sid), exist_ok=True)
        va = [u for t, u in dist if re.search(r"volume[ _]?a\s*\.xlsx$", t, re.I)]
        vc = [u for t, u in dist if re.search(r"volume ?c", t, re.I)]
        rec = {"dataset": ds, "issued": d.get("issued"), "title": (d.get("title") or {}).get("en"), "files": {}}
        if not va:
            rec["error"] = "δεν βρέθηκε Volume A"
            man["surveys"][sid] = rec
            print(sid, rec["error"], flush=True)
            continue
        pa = os.path.join(raw, sid, "volumeA.xlsx")
        if not os.path.exists(pa):
            curl(va[0], pa)
        rec["files"]["volumeA"] = {"url": va[0], "sha256": sha(open(pa, "rb").read())}
        items = parse_volume_a(pa)
        vis, lock = [], []
        for k, it in enumerate(items):
            iid = f"{sid}::{it['sheet']}"
            el = it["values"].pop(GREEK, None)
            nel = (it.get("nUnweighted") or {}).pop(GREEK, None)
            vis.append({"id": iid, "question": it["question"], "base": it["base"], "options": it["options"], "values": it["values"], "nUnweighted": it.get("nUnweighted")})
            lock.append({"id": iid, "EL": el, "nEL": nel})
        vj = json.dumps({"survey": sid, "items": vis}, ensure_ascii=False).encode()
        lj = json.dumps({"survey": sid, "items": lock}, ensure_ascii=False).encode()
        open(os.path.join(out, "visible", f"{sid}.json"), "wb").write(vj)
        open(os.path.join(out, "locked", f"{sid}.EL.json"), "wb").write(lj)
        rec.update({"nItems": len(vis), "visibleSha256": sha(vj), "lockedSha256": sha(lj), "greekPresent": sum(1 for x in lock if x["EL"] is not None)})
        if vc:
            pz = os.path.join(raw, sid, "volumeC.zip")
            try:
                if not os.path.exists(pz):
                    curl(vc[0], pz)
                z = zipfile.ZipFile(pz)
                gr = [n for n in z.namelist() if re.search(r"(^|[_\-\s/(])(EL|GR|Greece)([_\-\s.)]|$)", n)]
                if gr:
                    b = z.read(gr[0])
                    open(os.path.join(out, "locked", f"{sid}.volumeC.EL{os.path.splitext(gr[0])[1]}"), "wb").write(b)
                    rec["files"]["volumeC_EL"] = {"name": gr[0], "sha256": sha(b)}
                rec["files"]["volumeC"] = {"url": vc[0], "sha256": sha(open(pz, "rb").read()), "members": len(z.namelist())}
            except Exception as ex:
                rec["files"]["volumeC_error"] = repr(ex)[:200]
        man["surveys"][sid] = rec
        print(sid, "ερωτήσεις", len(vis), "με ελληνική στήλη", rec["greekPresent"], flush=True)
    json.dump(man, open(os.path.join(out, "manifest.json"), "w"), ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()
